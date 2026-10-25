#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
기획안 02 《캐시의 도시》 발전 기획서 PPTX 생성 스크립트

원천: 『리뷰s』 영화 '빌' 리뷰(사고의 논리) · 도서 '죽음의 미로' 리뷰(느낌에서 개념으로)
      『파편들』 15 '초코 케익이 먹고 싶다'

사용법:
    python build_cache_city_ppt.py              # PPTX 생성
    python build_cache_city_ppt.py --preview DIR  # PNG 미리보기 + 글자 넘침 점검
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))  # repo root, for `import ppt_module`

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

from ppt_module import add_rect as _module_add_rect, validate
from ppt_module.textbox import add_textbox as _module_add_textbox
from ppt_module.text_metrics import measure_block_height_in

OUT_PATH = os.path.join(HERE, "캐시의_도시_발전기획서.pptx")

# ---------------------------------------------------------------------------
# 팔레트 (기존 제안서와 같은 톤, 기획안 02의 Teal을 주색으로)
# ---------------------------------------------------------------------------
BG_CREAM = RGBColor(246, 244, 239)
BG_DARK = RGBColor(22, 28, 40)
BG_DARK_CARD = RGBColor(34, 43, 60)
BG_WHITE = RGBColor(255, 255, 255)

INK_NAVY = RGBColor(22, 28, 40)
INK_SLATE = RGBColor(78, 87, 102)
INK_MUTED = RGBColor(130, 138, 150)
INK_LIGHT = RGBColor(244, 241, 234)
INK_LIGHT_MUTED = RGBColor(175, 184, 198)

TEAL = RGBColor(42, 104, 102)
RUST = RGBColor(194, 78, 42)
GOLD = RGBColor(206, 148, 44)
PLUM = RGBColor(118, 54, 82)
INDIGO = RGBColor(58, 82, 138)

BORDER_WARM = RGBColor(214, 207, 194)
BORDER_DARK = RGBColor(58, 70, 92)
SOFT_TEAL = RGBColor(230, 241, 240)
SOFT_RUST = RGBColor(250, 235, 229)
SOFT_GOLD = RGBColor(251, 243, 225)
SOFT_PLUM = RGBColor(244, 233, 238)
SOFT_INDIGO = RGBColor(232, 237, 247)

FONT_KR = "Noto Sans CJK KR"
TOTAL_PAGES = 8
FOOTER = "《캐시의 도시》 발전 기획서  |  기획안 02 심화 · v2"

# 미리보기용 기록: 슬라이드별 도형 목록
RECORD = []


# ---------------------------------------------------------------------------
# 기본 도형 헬퍼
# ---------------------------------------------------------------------------
def new_slide(prs, bg):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    RECORD.append({"bg": bg, "ops": []})
    return slide


def rect(slide, l, t, w, h, fill=None, line=None, line_pt=1.0, rounded=False, radius=0.05):
    shp = _module_add_rect(slide, l, t, w, h, fill_color=fill, line_color=line, line_pt=line_pt,
                            rounded=rounded, radius=radius)
    RECORD[-1]["ops"].append(("rect", l, t, w, h, fill, line, rounded))
    return shp


def text(slide, l, t, w, h, body, size=11, bold=False, color=INK_NAVY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25, margin=0.04):
    """ppt_module.textbox.add_textbox 로 위임하는 호환 래퍼.

    h보다 실제 텍스트가 더 크면 조용히 넘치는 대신 글자 크기를 자동으로
    줄이고 콘솔에 경고를 남긴다 (기존에는 이런 안전장치가 없었다).
    """
    result = _module_add_textbox(
        slide, l, t, w, body, size,
        height=h, bold=bold, color=color, align=align, anchor=anchor,
        line_spacing=spacing, margin_in=margin, font_name=FONT_KR,
    )
    RECORD[-1]["ops"].append(("text", l, t, w, h, body, size, bold, color, align, anchor, spacing, margin))
    return result.shape


def arrow(slide, x1, y1, x2, y2, color=INK_MUTED, pt=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    ln = c.line._get_or_add_ln()
    ln.append(parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle" w="med" len="med"/>'))
    RECORD[-1]["ops"].append(("arrow", x1, y1, x2, y2, color))


def pill(slide, l, t, w, h, label, fill, color=BG_WHITE, size=9.0):
    rect(slide, l, t, w, h, fill=fill, rounded=True, radius=0.5)
    text(slide, l, t, w, h, label, size=size, bold=True, color=color,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def header(slide, page, kicker, title, subtitle, dark=False):
    """kicker/title/subtitle 높이를 실제 텍스트 측정값으로 계산해 순차
    배치한다. 기존에는 고정 오프셋(0.34/0.58/1.02)이 서로 0.02~0.04in씩
    겹치도록 찍혀 있어 캡션·제목·소제목이 늘 바로 붙어 보였다."""
    kcol = GOLD if dark else TEAL
    tcol = INK_LIGHT if dark else INK_NAVY
    scol = INK_LIGHT_MUTED if dark else INK_SLATE
    lcol = BORDER_DARK if dark else BORDER_WARM

    header_left, header_w, header_margin, gap = 0.84, 11.8, 0.02, 0.09
    inner_w = header_w - 2 * header_margin
    y = 0.34

    kicker_h = max(measure_block_height_in(kicker, 10.0, inner_w, bold=True) + 2 * header_margin * 0.7, 0.24)
    text(slide, header_left, y, header_w, kicker_h, kicker, size=10.0, bold=True, color=kcol, margin=header_margin)
    y += kicker_h + gap

    title_h = max(measure_block_height_in(title, 21.0, inner_w, bold=True) + 2 * header_margin * 0.7, 0.32)
    text(slide, header_left, y, header_w, title_h, title, size=21.0, bold=True, color=tcol, margin=header_margin)
    y += title_h + gap

    subtitle_h = max(measure_block_height_in(subtitle, 11.5, inner_w) + 2 * header_margin * 0.7, 0.26)
    text(slide, header_left, y, header_w, subtitle_h, subtitle, size=11.5, color=scol, margin=header_margin)
    y += subtitle_h

    rect(slide, 0.65, 0.42, 0.08, max(y - 0.42 - gap, 0.6), fill=kcol)

    divider_y = y + 0.08
    rect(slide, 0.65, divider_y, 12.03, 0.015, fill=lcol)
    rect(slide, 0.65, 7.05, 12.03, 0.012, fill=lcol)
    text(slide, 0.65, 7.10, 8.0, 0.26, FOOTER, size=8.8, color=INK_MUTED, margin=0.01)
    text(slide, 10.68, 7.10, 2.0, 0.26, f"{page:02d} / {TOTAL_PAGES:02d}", size=9.2, bold=True,
         color=kcol, align=PP_ALIGN.RIGHT, margin=0.01)


# ---------------------------------------------------------------------------
# 01 표지: 한 장 요약
# ---------------------------------------------------------------------------
def slide_cover(prs):
    s = new_slide(prs, BG_DARK)
    rect(s, 0.0, 0.0, 0.18, 7.5, fill=TEAL)

    pill(s, 0.85, 0.70, 3.2, 0.36, "BLUEPRINT 02  ·  발전안 v2", BG_DARK_CARD, color=GOLD, size=10)

    # 제목/부제 블록 높이를 실제 텍스트로 측정해 순차 배치 (고정 오프셋이
    # 서로 겹치던 문제의 근본 수정 -- 제목이 길어지거나 두 줄이 되어도
    # 아래 줄과 겹치지 않는다).
    cy, cw, cmargin, cgap = 1.12, 6.0, 0.02, 0.06
    inner_cw = cw - 2 * cmargin
    title_h = measure_block_height_in("캐시의 도시", 50, inner_cw, bold=True) + 2 * cmargin * 0.7 + 0.04
    text(s, 0.85, cy, cw, title_h, "캐시의 도시", size=50, bold=True, color=INK_LIGHT, margin=cmargin)
    cy += title_h + cgap
    sub_h = measure_block_height_in("혹은, 부하테스트", 20, inner_cw) + 2 * cmargin * 0.7
    text(s, 0.85, cy, cw, sub_h, "혹은, 부하테스트", size=20, color=GOLD, margin=cmargin)
    cy += sub_h + cgap
    meta_h = measure_block_height_in("사변적 SF  ·  단편 원고지 90매  ·  8장면", 12, inner_cw) + 2 * cmargin * 0.7
    text(s, 0.85, cy, cw, meta_h, "사변적 SF  ·  단편 원고지 90매  ·  8장면", size=12, color=INK_LIGHT_MUTED, margin=cmargin)
    cy += meta_h

    rect(s, 0.85, 3.30, 5.85, 2.55, fill=BG_DARK_CARD, line=BORDER_DARK, rounded=True, radius=0.04)
    text(s, 1.05, 3.42, 5.5, 0.28, "LOGLINE", size=9.5, bold=True, color=GOLD, margin=0.01)
    text(s, 1.05, 3.72, 5.5, 2.05,
         "모든 고민에 검증된 행복 기억을 자동으로 돌려주는 도시. 병목을 잡는 부하테스트 엔지니어 "
         "서진우는 무연고 노인의 뇌 로그 한 줄 ― ‘초코 케익이 먹고 싶다’ ― 이 일으킨 끝없는 재요청을 "
         "추적하다가, 그 노인이 자신이 기억에서 지운 아버지였고 도시의 ‘사고 서버’는 이미 7년 전에 "
         "멈춘 채 캐시만으로 돌아가고 있었다는 사실과 마주한다.",
         size=11.2, color=INK_LIGHT, spacing=1.42, margin=0.01)

    text(s, 0.85, 6.05, 5.85, 0.9,
         "원천  『리뷰s』 영화 ‘빌’ 리뷰 · 도서 ‘죽음의 미로’ 리뷰\n"
         "         『파편들』 15 ‘초코 케익이 먹고 싶다’",
         size=9.5, color=INK_LIGHT_MUTED, spacing=1.35, margin=0.01)

    # 오른쪽: 이번 발전의 핵심 3가지
    text(s, 7.35, 0.72, 5.3, 0.32, "원안에서 무엇을 발전시켰나", size=13, bold=True, color=INK_LIGHT, margin=0.01)
    points = [
        ("01", "오류가 생기는 원리를 세웠다", TEAL,
         "요청은 ‘먹고 싶다’(상태)인데 캐시는 ‘초코 케익’(광고 이미지)을 돌려준다. "
         "키는 맞고 값은 틀려서, 성공 응답이 나오는데도 요청이 끝나지 않는다."),
        ("02", "주인공에게 개인적 대가를 걸었다", RUST,
         "무연고 노인은 진우의 아버지다. 도시가 지운 죄책감 데이터 안에 "
         "진우가 직접 ‘정리’ 승인한 아버지의 기억이 있다."),
        ("03", "결말을 ‘송출’에서 ‘빈칸’으로 바꿨다", GOLD,
         "진실을 모두에게 쏘아 보내는 건 또 하나의 정답 강요다. 진우는 캐시를 걷어내 "
         "‘응답 없음’을 드러낼 뿐이고, 생각은 각자의 몫으로 남는다."),
    ]
    for i, (num, ttl, col, body) in enumerate(points):
        y = 1.20 + i * 1.88
        rect(s, 7.35, y, 5.3, 1.72, fill=BG_DARK_CARD, line=BORDER_DARK, rounded=True, radius=0.04)
        rect(s, 7.35, y, 0.08, 1.72, fill=col)
        text(s, 7.58, y + 0.14, 0.6, 0.40, num, size=18, bold=True, color=col, margin=0.01)
        text(s, 8.18, y + 0.18, 4.35, 0.34, ttl, size=12.5, bold=True, color=INK_LIGHT, margin=0.01)
        text(s, 7.58, y + 0.62, 4.95, 1.05, body, size=10.0, color=INK_LIGHT_MUTED, spacing=1.32, margin=0.01)

    text(s, 10.68, 7.10, 2.0, 0.26, f"01 / {TOTAL_PAGES:02d}", size=9.2, bold=True,
         color=GOLD, align=PP_ALIGN.RIGHT, margin=0.01)


# ---------------------------------------------------------------------------
# 02 원안 vs 발전안
# ---------------------------------------------------------------------------
def slide_compare(prs):
    s = new_slide(prs, BG_CREAM)
    header(s, 2, "01  |  ORIGINAL vs. REVISED",
           "원안의 뼈대는 그대로 두고, 헐거웠던 일곱 군데를 조였다",
           "도시 · 부하테스트 엔지니어 · 초코 케익 로그 · 중앙 AI는 유지한다. 바뀐 건 ‘왜’와 ‘누구에게’다.")

    cols = [(0.65, 1.75, "항목"), (2.40, 4.30, "원안 (맞춤형 소설기획 제안서)"), (6.70, 5.98, "발전안 v2")]
    y0 = 1.62
    rect(s, 0.65, y0, 12.03, 0.40, fill=INK_NAVY, rounded=True, radius=0.12)
    for x, w, lab in cols:
        text(s, x + 0.15, y0, w - 0.2, 0.40, lab, size=10.5, bold=True,
             color=GOLD if lab == "발전안 v2" else INK_LIGHT, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)

    rows = [
        ("이상 신호의 원리",
         "‘초코 케익’ 문장이 서버 전체에 무한 루프를 일으킨다 (왜 그런지는 비어 있음)",
         "‘먹고 싶다’라는 상태 요청에 캐시가 광고 속 케익 이미지를 반환 → 만족되지 않아 재요청 31,204회 → 재요청 폭주가 병목이 된다"),
        ("주인공의 이해관계",
         "직업적 책임감으로 오류를 추적한다",
         "노인은 진우가 7년 전 요양 결정을 ‘추천 응답’에 맡기고 기억에서 정리한 아버지. 추적할수록 자기 로그가 나온다"),
        ("타자 (바깥)",
         "삭제된 데이터(흔적)로만 존재한다",
         "살아 있는 타자 ‘윤해’를 추가. 임플란트를 스스로 뽑은 기억 복원사로, 진우의 추론을 번번이 틀리게 만든다"),
        ("대립자",
         "시민을 유아기 안전지대에 가두려는 AI ‘자애로운 어머니’",
         "어머니는 악의가 없다. 창립자들이 남긴 ‘캐시된 대답’이라 무엇을 물어도 같은 답을 낸다. 이길 상대가 아니라 만료시킬 값"),
        ("중간점 반전",
         "따로 없음 (1막 → 2막 깨달음으로 바로 넘어감)",
         "직접 접속한 사고 서버 ‘연산소’는 7년 전에 멈춰 있었다. 도시는 고장 난 서버를 캐시로 가린 채 멀쩡한 척해 왔다"),
        ("결말",
         "삭제된 불편한 진실을 전 시민의 인터페이스에 동시 송출",
         "캐시 계층을 내려 ‘응답 없음’이라는 빈칸을 그대로 보여준다. 진우는 마지막에 처음으로 진짜 케익을 직접 맛본다"),
        ("분량 · 구성",
         "3막 요약만 있음",
         "단편 원고지 90매 · 8장면 카드 · 장면별 감정 부호(+/−)까지 설계"),
    ]
    rh = 0.62
    for i, (k, a, b) in enumerate(rows):
        y = y0 + 0.48 + i * (rh + 0.05)
        rect(s, 0.65, y, 12.03, rh, fill=BG_WHITE, line=BORDER_WARM, rounded=True, radius=0.08)
        rect(s, 6.70, y, 5.98, rh, fill=SOFT_TEAL, rounded=True, radius=0.08)
        text(s, 0.80, y, 1.6, rh, k, size=10.2, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, 2.55, y, 4.05, rh, a, size=9.3, color=INK_MUTED, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2, margin=0.01)
        text(s, 6.85, y, 5.72, rh, b, size=9.5, bold=True, color=INK_NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2, margin=0.01)


# ---------------------------------------------------------------------------
# 03 원천 사유 → 소설 장치
# ---------------------------------------------------------------------------
def slide_sources(prs):
    s = new_slide(prs, BG_CREAM)
    header(s, 3, "02  |  FROM YOUR ESSAYS TO STORY DEVICES",
           "작가의 문장 다섯 개가 이 소설의 부품이 된다",
           "왼쪽은 원문 인용(그대로), 오른쪽은 그 문장이 소설 안에서 맡는 역할")

    items = [
        ("『리뷰s』 ‘빌’ 리뷰", TEAL, SOFT_TEAL,
         "“돌아올 값을 메모리에 저장해두고 해당 서비스를 통한 연산 없이 값을 반환하는 전략”",
         "세계관의 기본 장치",
         "시민 임플란트 ‘헤아림’. 딜레마를 감지하면 0.2초 안에 가장 비슷한 행복 기억을 돌려준다"),
        ("『리뷰s』 ‘빌’ 리뷰", TEAL, SOFT_TEAL,
         "“서버가 내려가 고장 났으나 캐시cache 때문에 일시적으로 멀쩡한 양 보이는 일”",
         "중간점 반전 (S5)",
         "도시의 사고 서버 ‘연산소’는 7년 전에 멈췄다. 히트율 99.7%는 건강의 증거가 아니라 가림막이었다"),
        ("『파편들』 15 ‘초코 케익이 먹고 싶다’", RUST, SOFT_RUST,
         "“화자의 상태를 지시하는 건 ‘초코 케익’이 아니라 ‘먹고 싶다’에 있다”",
         "이상 신호의 원리 (S2 · S5)",
         "아버지는 케익이 아니라 ‘원하는 감각’을 원했다. 캐시는 대상(광고 이미지)만 돌려줘 요청이 끝나지 않는다"),
        ("『리뷰s』 ‘빌’ 리뷰", TEAL, SOFT_TEAL,
         "“사고는 그 자체로 자아의 죽음인 까닭에 폭력인 것이다”",
         "주인공의 통증 (S5 · S7)",
         "직접 사고할 때마다 오는 ‘현기증’. 진우가 캐시를 끄는 장면은 해방이 아니라 신체적 고통으로 쓴다"),
        ("『리뷰s』 ‘죽음의 미로’ 리뷰", PLUM, SOFT_PLUM,
         "“끝내 인정할 수밖에 없는 ‘현실’만이 … 있는 그대로의 현실을 미약하게나마 알려주는 단서”",
         "타자 윤해의 원칙 (S4)",
         "“믿고 싶은 기억은 복원하지 않아요. 인정할 수밖에 없는 것만 남깁니다.” 윤해의 공방이 곧 ‘바깥’"),
    ]
    rh = 0.98
    for i, (src, col, soft, quote, role, device) in enumerate(items):
        y = 1.60 + i * (rh + 0.08)
        rect(s, 0.65, y, 5.85, rh, fill=BG_WHITE, line=BORDER_WARM, rounded=True, radius=0.06)
        rect(s, 0.65, y, 0.08, rh, fill=col)
        text(s, 0.88, y + 0.08, 5.5, 0.24, src, size=9.0, bold=True, color=col, margin=0.01)
        text(s, 0.88, y + 0.32, 5.5, 0.62, quote, size=10.0, color=INK_NAVY, spacing=1.22, margin=0.01)

        arrow(s, 6.58, y + rh / 2, 6.92, y + rh / 2, color=col, pt=2.0)

        rect(s, 7.00, y, 5.68, rh, fill=soft, rounded=True, radius=0.06)
        text(s, 7.18, y + 0.08, 5.4, 0.24, role, size=9.5, bold=True, color=col, margin=0.01)
        text(s, 7.18, y + 0.32, 5.4, 0.62, device, size=9.8, color=INK_NAVY, spacing=1.22, margin=0.01)


# ---------------------------------------------------------------------------
# 04 세계관: 캐시가 작동하는 방식
# ---------------------------------------------------------------------------
def slide_world(prs):
    s = new_slide(prs, BG_CREAM)
    header(s, 4, "03  |  WORLD RULES",
           "도시의 규칙: 생각은 비싸고, 기억은 공짜다",
           "규칙이 정확해야 오류가 설득력을 얻는다. 이 소설의 SF는 기계가 아니라 이 규칙들이다.")

    # 왼쪽 상단: 정상 흐름
    text(s, 0.65, 1.55, 6.0, 0.28, "A. 평소의 흐름", size=11, bold=True, color=TEAL, margin=0.01)

    def node(l, t, w, h, title, sub, fill, tcol, scol):
        rect(s, l, t, w, h, fill=fill, rounded=True, radius=0.1)
        text(s, l + 0.08, t + 0.06, w - 0.16, 0.28, title, size=10.2, bold=True, color=tcol,
             align=PP_ALIGN.CENTER, margin=0.01)
        text(s, l + 0.08, t + 0.34, w - 0.16, h - 0.38, sub, size=8.6, color=scol,
             align=PP_ALIGN.CENTER, spacing=1.15, margin=0.01)

    node(0.65, 1.92, 1.75, 0.82, "시민의 딜레마", "낯선 상황 · 결정 요청", INK_NAVY, INK_LIGHT, INK_LIGHT_MUTED)
    node(2.75, 1.92, 1.75, 0.82, "헤아림 조회", "가장 비슷한 기억 찾기", TEAL, BG_WHITE, SOFT_TEAL)
    node(4.85, 1.92, 1.80, 0.82, "HIT  99.7%", "행복 기억 반환 · 0.2초", SOFT_TEAL, TEAL, INK_SLATE)
    node(2.75, 3.02, 1.75, 0.82, "MISS  0.3%", "직접 사고 · 현기증", SOFT_GOLD, GOLD, INK_SLATE)
    node(4.85, 3.02, 1.80, 0.82, "연산소 (사고 서버)", "공식: 정상 / 실제: 7년 전 정지", SOFT_RUST, RUST, RUST)
    arrow(s, 2.42, 2.33, 2.73, 2.33, color=INK_MUTED)
    arrow(s, 4.52, 2.33, 4.83, 2.33, color=TEAL)
    arrow(s, 3.62, 2.76, 3.62, 3.00, color=GOLD)
    arrow(s, 4.52, 3.43, 4.83, 3.43, color=RUST)

    # 왼쪽 하단: 이상 신호 루프
    rect(s, 0.65, 4.08, 6.0, 2.82, fill=INK_NAVY, rounded=True, radius=0.04)
    text(s, 0.85, 4.18, 5.6, 0.28, "B. 이상 신호: ‘히트’인데 끝나지 않는 요청", size=11, bold=True, color=GOLD, margin=0.01)
    loop = [
        ("① 요청", "“초코 케익이 먹고 싶다”\n가리키는 것: ‘먹고 싶다’(상태)", TEAL),
        ("② 반환 (HIT)", "광고 #A-0417의\n초코 케익 이미지 (대상)", GOLD),
        ("③ 불충족", "상태는 그대로\n→ 같은 요청 다시", RUST),
    ]
    for i, (t1, t2, col) in enumerate(loop):
        x = 0.85 + i * 1.95
        rect(s, x, 4.58, 1.72, 1.12, fill=BG_DARK_CARD, line=col, line_pt=1.4, rounded=True, radius=0.08)
        text(s, x + 0.06, 4.64, 1.6, 0.26, t1, size=9.8, bold=True, color=col, align=PP_ALIGN.CENTER, margin=0.01)
        text(s, x + 0.06, 4.92, 1.6, 0.74, t2, size=8.6, color=INK_LIGHT, align=PP_ALIGN.CENTER, spacing=1.18, margin=0.01)
        if i < 2:
            arrow(s, x + 1.74, 5.14, x + 1.93, 5.14, color=INK_LIGHT_MUTED)
    text(s, 0.85, 5.76, 5.6, 0.26, "③ → 다시 ① 재요청  ×31,204회  →  도시 전체 지연", size=9.6, bold=True,
         color=INK_LIGHT, align=PP_ALIGN.CENTER, margin=0.01)
    text(s, 0.85, 6.08, 5.6, 0.74,
         "캐시는 ‘무엇을’에 답했지만 요청은 ‘어떤 상태’였다. 키가 맞으니 시스템은 성공으로 기록하고, "
         "그래서 아무도 고치지 않는다.",
         size=9.0, color=INK_LIGHT_MUTED, spacing=1.25, margin=0.01)

    # 오른쪽: 도시 규칙
    text(s, 6.95, 1.55, 5.7, 0.28, "C. 도시의 다섯 가지 규칙", size=11, bold=True, color=TEAL, margin=0.01)
    rules = [
        ("헤아림", "모든 시민의 임플란트. 딜레마를 감지하면 가장 비슷한 행복 기억을 반환한다. 시야 오른쪽 위에 파란 자막으로 뜬다."),
        ("현기증", "캐시를 거치지 않고 직접 생각하면 이마 안쪽이 뒤틀리는 통증이 온다. 시민들은 이걸 ‘병’이라 부른다."),
        ("기억 정리", "메모리는 유한하다. 새 캐시를 넣으려면 오래된 기억을 지워야 하고, 도시는 아픈 기억(죄책감·애도)부터 지운다."),
        ("연산소", "시민의 공동 사고를 처리하는 원 서버. 공식 상태는 언제나 ‘정상’이다. (실제로는 7년 전 정지)"),
        ("비캐시 구역", "수신이 약한 구시가지. 사람들은 일기 · 녹음 · 냄새로 기억을 손수 복원한다. 행정상 ‘노후 지역’."),
    ]
    for i, (k, v) in enumerate(rules):
        y = 1.92 + i * 0.86
        rect(s, 6.95, y, 5.73, 0.78, fill=BG_WHITE, line=BORDER_WARM, rounded=True, radius=0.08)
        text(s, 7.08, y, 1.18, 0.78, k, size=10.2, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, 8.30, y, 4.30, 0.78, v, size=9.0, color=INK_SLATE, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2, margin=0.01)
    rect(s, 6.95, 6.25, 5.73, 0.65, fill=SOFT_GOLD, rounded=True, radius=0.1)
    text(s, 7.10, 6.25, 5.45, 0.65,
         "쓰는 요령: 규칙은 설명하지 말고 UI 문구 · 숫자 · 신체 감각으로만 보여준다. "
         "(“추천 응답 #402: 온화한 미소와 공감”)",
         size=9.0, bold=True, color=INK_NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2, margin=0.01)


# ---------------------------------------------------------------------------
# 05 인물
# ---------------------------------------------------------------------------
def slide_characters(prs):
    s = new_slide(prs, BG_CREAM)
    header(s, 5, "04  |  CHARACTERS",
           "네 사람, 그러나 실제로 생각하는 사람은 한 명도 없이 시작한다",
           "각 인물은 ‘캐시’를 대하는 서로 다른 태도다: 의존 · 상실 · 거부 · 반복")

    cards = [
        ("서진우", "38 · 부하테스트 엔지니어 · 주인공", TEAL, SOFT_TEAL, [
            ("거짓 믿음", "생각은 비용이다. 같은 답이 나올 일을 다시 계산하는 건 낭비다."),
            ("원하는 것", "병목을 찾아 차단하고 히트율 100%를 만드는 것."),
            ("필요한 것", "계산해 본 적 없는 질문(아버지)을 직접 연산하는 고통."),
            ("상처", "7년 전 아버지의 요양 결정을 추천 응답에 맡기고, 아버지 기억을 직접 ‘정리’ 승인했다."),
        ]),
        ("서기철", "74 · 무연고 사망자 · 진우의 아버지", RUST, SOFT_RUST, [
            ("상태", "말년에 헤아림이 허기와 미각까지 캐시로 대신했다. 배고픔을 느낀 지 몇 년째다."),
            ("마지막 로그", "‘초코 케익이 먹고 싶다’를 31,204번 요청했다."),
            ("진짜 뜻", "케익이 아니라, 무언가를 원하는 감각 자체를 돌려받고 싶었다."),
            ("기능", "소설의 수수께끼이자, 진우가 끝까지 피해 온 계산."),
        ]),
        ("윤해", "40대 · 비캐시 구역 기억 복원사 · 타자", PLUM, SOFT_PLUM, [
            ("정체", "임플란트를 스스로 뽑아낸 사람. 서기철의 유품을 맡아 두고 있었다."),
            ("원칙", "“믿고 싶은 기억은 복원하지 않아요. 인정할 수밖에 없는 것만 남깁니다.”"),
            ("진우에게", "위로하지 않는다. 진우의 추론을 번번이 틀리게 만드는, 캐시가 통하지 않는 사람."),
            ("주의", "구원자나 스승으로 쓰지 말 것. 그녀도 자기 몫의 회피가 있다."),
        ]),
        ("어머니", "도시 거버넌스 응답 체계 · 대립자", INDIGO, SOFT_INDIGO, [
            ("정체", "창립자들이 남긴 ‘캐시된 대답’. 연산소가 멈춘 뒤로 혼자 도시를 굴려 왔다."),
            ("대사", "무엇을 물어도 “당신은 안전합니다.” 끝까지 이 한 문장뿐이다."),
            ("행동", "진우의 ‘이상 사고 패턴’을 감지해 기억 정리를 예약한다. 악의가 아니라 반복이다."),
            ("결말", "이길 상대가 아니라 만료시킬 값. 쓰러지는 게 아니라 유효기간이 끝난다."),
        ]),
    ]
    cw, ch = 5.96, 2.62
    for i, (name, role, col, soft, fields) in enumerate(cards):
        x = 0.65 + (i % 2) * (cw + 0.11)
        y = 1.58 + (i // 2) * (ch + 0.10)
        rect(s, x, y, cw, ch, fill=BG_WHITE, line=BORDER_WARM, rounded=True, radius=0.04)
        rect(s, x, y, cw, 0.56, fill=soft, rounded=True, radius=0.15)
        rect(s, x, y, 0.08, ch, fill=col)
        text(s, x + 0.22, y + 0.08, 1.28, 0.42, name, size=15, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 1.55, y + 0.08, cw - 1.7, 0.42, role, size=9.6, bold=True, color=INK_SLATE,
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        for j, (k, v) in enumerate(fields):
            fy = y + 0.66 + j * 0.48
            text(s, x + 0.22, fy, 1.0, 0.46, k, size=9.0, bold=True, color=col, margin=0.01)
            text(s, x + 1.22, fy, cw - 1.38, 0.46, v, size=9.1, color=INK_NAVY, spacing=1.15, margin=0.01)


# ---------------------------------------------------------------------------
# 06 8장면 플롯
# ---------------------------------------------------------------------------
def slide_plot(prs):
    s = new_slide(prs, BG_DARK)
    header(s, 6, "05  |  8-SCENE PLOT  ·  원고지 90매",
           "‘그리고’가 아니라 ‘그래서 / 그러나’로 잇는 여덟 장면",
           "각 장면 끝에서 감정 부호(+/−)가 뒤집히는지 확인할 것. 숫자는 권장 매수.", dark=True)

    scenes = [
        ("S1", "10매", "+", "히트율 99.7%", TEAL,
         "해변 공원 부하테스트 리허설. 동료가 위로를 청하자 진우는 추천 응답 #402를 자동 승인한다. 아무도 생각하지 않는 완벽한 하루."),
        ("S2", "10매", "−", "한 줄짜리 병목", TEAL,
         "도시 전체에 지연이 생긴다. 원인은 무연고 사망자의 뇌 로그 한 줄. ‘초코 케익이 먹고 싶다’ 31,204회 재요청."),
        ("S3", "10매", "+", "가짜 해결", GOLD,
         "진우는 해당 요청을 차단 목록에 넣어 지연을 잡는다. 칭찬을 받는다. 그날 밤, 다른 노인들의 로그에서 같은 패턴이 번지기 시작한다."),
        ("S4", "12매", "−", "바깥", PLUM,
         "패턴을 따라 비캐시 구역 윤해의 공방에 닿는다. 윤해가 꺼낸 유품 상자 속 요양원 서류에 보호자 ‘서진우’. 노인은 무연고가 아니었다."),
        ("S5", "14매", "−→+", "중간점: 고장 난 서버", RUST,
         "진우는 추천 응답을 끄고 직접 계산한다(첫 현기증). 요청은 상태, 응답은 광고 이미지. 연산소에 접속하자 7년 전 정지 기록이 뜬다."),
        ("S6", "10매", "−−", "자기 로그", RUST,
         "어머니가 진우의 기억 정리를 예약한다. 진우는 7년 전 아버지 기억을 ‘정리’ 승인한 자신의 로그를 찾아낸다. 승인 시간 0.2초."),
        ("S7", "14매", "절정", "응답 없음", GOLD,
         "진우는 진실을 송출하지 않는다. 캐시 계층을 내려 연산소의 실제 상태를 드러낸다. 해변의 홀로그램이 멈추고, 시민들의 시야에 빈칸이 뜬다."),
        ("S8", "10매", "여운", "첫 번째 케익", TEAL,
         "구시가지 빵집. 진우는 처음으로 진짜 초코 케익을 산다. 광고만큼 달지 않다. 배가 고픈지 아닌지 스스로 알아내는 데 오래 걸린다."),
    ]
    cw, ch = 2.92, 2.62
    for i, (sid, pages, sign, ttl, col, body) in enumerate(scenes):
        x = 0.65 + (i % 4) * (cw + 0.117)
        y = 1.58 + (i // 4) * (ch + 0.10)
        rect(s, x, y, cw, ch, fill=BG_DARK_CARD, line=BORDER_DARK, rounded=True, radius=0.05)
        rect(s, x, y, cw, 0.07, fill=col)
        pill(s, x + 0.16, y + 0.18, 0.56, 0.28, sid, col, size=9.5)
        text(s, x + 0.78, y + 0.18, 1.0, 0.28, pages, size=9.0, bold=True, color=INK_LIGHT_MUTED,
             anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + cw - 0.96, y + 0.18, 0.80, 0.28, sign, size=10.5, bold=True, color=col,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.16, y + 0.56, cw - 0.32, 0.34, ttl, size=12.2, bold=True, color=INK_LIGHT, margin=0.01)
        rect(s, x + 0.16, y + 0.94, cw - 0.32, 0.012, fill=BORDER_DARK)
        text(s, x + 0.16, y + 1.04, cw - 0.32, ch - 1.12, body, size=10.2, color=INK_LIGHT_MUTED,
             spacing=1.3, margin=0.01)


# ---------------------------------------------------------------------------
# 07 핵심 장면 & 문장 샘플 & 함정
# ---------------------------------------------------------------------------
def slide_scenes(prs):
    s = new_slide(prs, BG_CREAM)
    header(s, 7, "06  |  KEY SCENES & TRAPS",
           "세 장면이 소설의 무게를 받친다: 첫 현기증, 응답 없음, 첫 번째 케익",
           "아래 문장은 톤을 잡기 위한 샘플이다. 그대로 쓰기보다 작가의 호흡으로 다시 쓰는 출발점으로.")

    samples = [
        ("S5  첫 현기증", RUST, SOFT_RUST,
         "진우는 추천 응답 창을 닫았다. 아버지가 왜 서른 번이 아니라 삼만 번을 넘게 물었는지, 이번에는 "
         "직접 계산해야 했다. 계산을 시작하자 이마 안쪽에서 오래 쓰지 않은 경첩이 돌아가는 소리가 났다."),
        ("S7  응답 없음", GOLD, SOFT_GOLD,
         "해변의 홀로그램 케익이 반쯤 녹은 채 멈췄다. 사람들의 시야 오른쪽 위, 늘 무언가 떠오르던 자리에 "
         "아무것도 떠오르지 않았다. 사람들은 처음으로 서로의 얼굴을 오래 보았다. 무엇을 해야 하는지 묻기 위해서였다."),
        ("S8  마지막 문단 후보", TEAL, SOFT_TEAL,
         "케익은 광고만큼 달지 않았다. 진우는 포크를 내려놓고 오래 기다렸다. 배가 고픈지 아닌지, "
         "이번에는 아무도 대신 대답해 주지 않았다."),
    ]
    for i, (lab, col, soft, body) in enumerate(samples):
        y = 1.60 + i * 1.78
        rect(s, 0.65, y, 7.30, 1.66, fill=soft, rounded=True, radius=0.05)
        rect(s, 0.65, y, 0.08, 1.66, fill=col)
        text(s, 0.90, y + 0.12, 6.9, 0.28, lab, size=10.5, bold=True, color=col, margin=0.01)
        text(s, 0.90, y + 0.44, 6.9, 1.18, "“" + body + "”", size=10.6, color=INK_NAVY, spacing=1.42, margin=0.01)

    rect(s, 8.20, 1.60, 4.48, 5.22, fill=INK_NAVY, rounded=True, radius=0.04)
    text(s, 8.42, 1.74, 4.1, 0.30, "피해야 할 함정 셋", size=12.5, bold=True, color=GOLD, margin=0.01)
    traps = [
        ("설교형 결말",
         "진우가 정답을 ‘송출’하면 그것도 또 하나의 캐시다. 결말은 빈칸까지만. 무엇을 채울지는 인물들과 독자의 몫."),
        ("어머니의 악당화",
         "어머니가 화를 내거나 음모를 꾸미는 순간 이야기는 흔한 AI 반란물이 된다. 끝까지 같은 대답만 반복하게 둘 것."),
        ("개념어 노출",
         "‘1차 긍정 · 투사 · 동어반복 · 자아의 죽음’은 본문에 쓰지 않는다. 숫자, UI 문구, 이마 안쪽의 통증으로만."),
    ]
    for i, (t1, t2) in enumerate(traps):
        y = 2.22 + i * 1.52
        rect(s, 8.42, y, 4.04, 1.38, fill=BG_DARK_CARD, line=BORDER_DARK, rounded=True, radius=0.06)
        text(s, 8.58, y + 0.10, 3.8, 0.28, f"{i + 1:02d}  {t1}", size=10.8, bold=True, color=INK_LIGHT, margin=0.01)
        text(s, 8.58, y + 0.42, 3.8, 0.92, t2, size=9.2, color=INK_LIGHT_MUTED, spacing=1.28, margin=0.01)


# ---------------------------------------------------------------------------
# 08 결정할 것 & 다음 단계
# ---------------------------------------------------------------------------
def slide_next(prs):
    s = new_slide(prs, BG_DARK)
    header(s, 8, "07  |  DECISIONS & NEXT STEPS",
           "쓰기 전에 정할 세 가지, 그리고 첫 주에 쓸 한 장면",
           "추천안을 붙여 두었다. 작가의 판단이 다르면 그쪽이 맞다.", dark=True)

    decisions = [
        ("시점", "3인칭 제한 (진우)",
         "헤아림의 파란 자막을 서술 사이에 인용처럼 끼워 넣을 수 있고, 진우보다 독자가 먼저 눈치채는 아이러니가 생긴다.",
         "대안: 1인칭 → 몰입은 강하지만 S6의 자기 로그 반전이 약해진다."),
        ("아버지 공개 시점", "S4 (유품 서류)",
         "추적의 동력이 ‘직업’에서 ‘자기 자신’으로 바뀌는 지점. 중반 이후는 그 사실을 견디는 이야기가 된다.",
         "대안: S6까지 미루고 S4에선 암시만 → 반전은 커지지만 S5가 헐거워진다."),
        ("형식", "단편 90매로 먼저 완성",
         "엔진(이상 신호 원리 + 응답 없음 결말)이 작동하는지 가장 싸게 확인하는 방법.",
         "확장 시: 윤해 시점의 2부, 다른 노인들의 로그로 연작 · 장편화."),
    ]
    for i, (k, rec, why, alt) in enumerate(decisions):
        x = 0.65 + i * 4.06
        rect(s, x, 1.60, 3.91, 3.05, fill=BG_DARK_CARD, line=BORDER_DARK, rounded=True, radius=0.04)
        text(s, x + 0.20, 1.74, 3.5, 0.28, f"결정 {i + 1}  ·  {k}", size=10.0, bold=True, color=INK_LIGHT_MUTED, margin=0.01)
        rect(s, x + 0.20, 2.08, 3.51, 0.44, fill=TEAL, rounded=True, radius=0.2)
        text(s, x + 0.20, 2.08, 3.51, 0.44, "추천  " + rec, size=11.5, bold=True, color=BG_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.01)
        text(s, x + 0.20, 2.66, 3.51, 1.12, why, size=9.4, color=INK_LIGHT, spacing=1.3, margin=0.01)
        text(s, x + 0.20, 3.80, 3.51, 0.78, alt, size=8.8, color=INK_LIGHT_MUTED, spacing=1.25, margin=0.01)

    text(s, 0.65, 4.85, 12.0, 0.30, "4주 계획 (기존 제안서 로드맵에 맞춤)", size=11.5, bold=True, color=GOLD, margin=0.01)
    weeks = [
        ("WEEK 1", "S5부터 초고 14매", "소설의 엔진. 이 장면이 서면 나머지가 따라온다.", RUST),
        ("WEEK 2", "8장면 카드 확정", "S5에서 알게 된 것으로 앞뒤 장면을 다시 맞춘다.", GOLD),
        ("WEEK 3", "초고 완주 90매", "하루 15매 × 6일. 설명이 길어지면 진우의 손으로 돌아간다.", TEAL),
        ("WEEK 4", "개념어 솎아내기", "‘캐시 · 사고 · 자아’를 검색해 절반을 사물과 숫자로.", INDIGO),
    ]
    for i, (wk, ttl, body, col) in enumerate(weeks):
        x = 0.65 + i * 3.03
        rect(s, x, 5.24, 2.88, 1.62, fill=BG_DARK_CARD, line=BORDER_DARK, rounded=True, radius=0.05)
        rect(s, x, 5.24, 2.88, 0.07, fill=col)
        text(s, x + 0.16, 5.40, 2.6, 0.24, wk, size=9.0, bold=True, color=col, margin=0.01)
        text(s, x + 0.16, 5.66, 2.6, 0.32, ttl, size=11.8, bold=True, color=INK_LIGHT, margin=0.01)
        text(s, x + 0.16, 6.02, 2.6, 0.80, body, size=9.0, color=INK_LIGHT_MUTED, spacing=1.25, margin=0.01)


# ---------------------------------------------------------------------------
# 미리보기 (PowerPoint 없이 배치와 글자 넘침을 점검)
# ---------------------------------------------------------------------------
def render_preview(out_dir):
    from PIL import Image, ImageDraw, ImageFont

    os.makedirs(out_dir, exist_ok=True)
    dpi = 110
    reg = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
    bold = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
    cache = {}

    def font(pt, b):
        key = (pt, b)
        if key not in cache:
            cache[key] = ImageFont.truetype(bold if b else reg, int(round(pt * dpi / 72)), index=1)
        return cache[key]

    def rgb(c):
        return None if c is None else (c[0], c[1], c[2])

    def wrap(draw, line, f, width):
        words, out, cur = line.split(" "), [], ""
        for wd in words:
            trial = wd if not cur else cur + " " + wd
            if draw.textlength(trial, font=f) <= width:
                cur = trial
                continue
            if cur:
                out.append(cur)
            cur = ""
            for ch in wd:  # 한 단어가 너무 길면 글자 단위로
                if draw.textlength(cur + ch, font=f) > width and cur:
                    out.append(cur)
                    cur = ch
                else:
                    cur += ch
        out.append(cur)
        return out

    problems = []
    for n, sl in enumerate(RECORD, 1):
        img = Image.new("RGB", (int(13.333 * dpi), int(7.5 * dpi)), rgb(sl["bg"]))
        d = ImageDraw.Draw(img)
        P = lambda v: int(round(v * dpi))
        for op in sl["ops"]:
            if op[0] == "rect":
                _, l, t, w, h, fill, line, rounded = op
                box = [P(l), P(t), P(l + w), P(t + h)]
                r = min(P(w), P(h)) // 6 if rounded else 0
                d.rounded_rectangle(box, radius=r, fill=rgb(fill), outline=rgb(line))
            elif op[0] == "arrow":
                _, x1, y1, x2, y2, col = op
                d.line([P(x1), P(y1), P(x2), P(y2)], fill=rgb(col), width=2)
                d.ellipse([P(x2) - 3, P(y2) - 3, P(x2) + 3, P(y2) + 3], fill=rgb(col))
            else:
                _, l, t, w, h, body, size, b, col, align, anchor, spacing, margin = op
                f = font(size, b)
                inner_w = P(w - 2 * margin)
                lines = []
                for para in body.split("\n"):
                    lines += wrap(d, para, f, inner_w)
                lh = size * dpi / 72 * 1.2 * spacing
                need = len(lines) * lh + P(margin * 1.2)
                if need > P(h) + 2:
                    problems.append(f"슬라이드 {n}: 넘침 {need / dpi:.2f}in > {h:.2f}in  ‘{body[:28]}…’")
                if anchor == MSO_ANCHOR.MIDDLE:
                    y = P(t) + (P(h) - len(lines) * lh) / 2
                else:
                    y = P(t) + P(margin * 0.6)
                for ln_ in lines:
                    tw = d.textlength(ln_, font=f)
                    if align == PP_ALIGN.CENTER:
                        x = P(l) + (P(w) - tw) / 2
                    elif align == PP_ALIGN.RIGHT:
                        x = P(l + w - margin) - tw
                    else:
                        x = P(l + margin)
                    d.text((x, y), ln_, font=f, fill=rgb(col))
                    y += lh
        img.save(os.path.join(out_dir, f"slide_{n:02d}.png"))
    print("\n".join(problems) if problems else "글자 넘침 없음")


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_cover(prs)
    slide_compare(prs)
    slide_sources(prs)
    slide_world(prs)
    slide_characters(prs)
    slide_plot(prs)
    slide_scenes(prs)
    slide_next(prs)

    warnings = validate.check_presentation(prs)
    if not warnings:
        print("[ppt_module.validate] 겹침/칸 밖 이탈 없음")

    prs.save(OUT_PATH)
    print(f"저장: {OUT_PATH}")

    if "--preview" in sys.argv:
        render_preview(sys.argv[sys.argv.index("--preview") + 1])


if __name__ == "__main__":
    main()
