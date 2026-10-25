#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
소설 쓰기 마스터클래스 (The Craft of Fiction Writing) PPTX 생성 스크립트
"""

import os
import math
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from ppt_module import set_slide_background, add_rect, validate
from ppt_module.textbox import add_textbox as _module_add_textbox
from ppt_module.text_metrics import measure_block_height_in

# ---------------------------------------------------------------------------
# Color Palette (Literary Atelier)
# ---------------------------------------------------------------------------
BG_CREAM = RGBColor(248, 245, 240)       # #F8F5F0 Warm Paper
BG_DARK = RGBColor(26, 33, 46)           # #1A212E Deep Ink Navy
BG_DARK_CARD = RGBColor(38, 47, 64)      # #262F40 Elevated Dark Card
BG_SAND = RGBColor(238, 233, 224)        # #EEE9E0 Warm Sand Card
BG_WHITE = RGBColor(255, 255, 255)       # #FFFFFF Crisp Card

INK_NAVY = RGBColor(26, 33, 46)          # #1A212E Primary Dark Text
INK_SLATE = RGBColor(85, 94, 108)        # #555E6C Secondary Body Text
INK_MUTED = RGBColor(135, 142, 153)      # #878E99 Muted Caption / Footer
INK_LIGHT = RGBColor(242, 239, 233)      # #F2EFE9 Text on Dark BG
INK_LIGHT_MUTED = RGBColor(178, 186, 198)# #B2BAC6 Subtitle on Dark BG

ACCENT_RUST = RGBColor(198, 83, 45)      # #C6532D Literary Terracotta
ACCENT_GOLD = RGBColor(209, 152, 48)     # #D19830 Antique Ochre Gold
ACCENT_TEAL = RGBColor(49, 107, 105)     # #316B69 Sage Teal
ACCENT_BURGUNDY = RGBColor(142, 53, 66)  # #8E3542 Deep Rose/Burgundy

BORDER_WARM = RGBColor(218, 211, 199)    # #DAD3C7 Subtle Card Border
BORDER_DARK = RGBColor(62, 74, 96)       # #3E4A60 Dark Card Border
SOFT_RUST_BG = RGBColor(250, 236, 230)   # #FAECE6 Highlight Tint
SOFT_TEAL_BG = RGBColor(232, 242, 241)   # #E8F2F1 Teal Tint
SOFT_GOLD_BG = RGBColor(251, 244, 228)   # #FBF4E4 Gold Tint

FONT_KR = "Noto Sans CJK KR"

ASSET_DIR = "/tmp/ppt_novel_assets"
os.makedirs(ASSET_DIR, exist_ok=True)


# ---------------------------------------------------------------------------
# Pillow High-Res Visual Asset Generators
# ---------------------------------------------------------------------------
def get_pil_font(size, bold=False):
    candidates = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc" if bold else "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Medium.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                pass
    return ImageFont.load_default()


def generate_cover_art():
    """Generates an editorial illustration for the cover slide (dark navy theme)."""
    w, h = 1150, 920
    img = Image.new("RGBA", (w, h), (26, 33, 46, 0))
    draw = ImageDraw.Draw(img)

    # Concentric editorial geometric rings (golden ratio / story cosmos)
    cx, cy = int(w * 0.54), int(h * 0.48)
    for r, col, width in [
        (410, (62, 74, 96, 95), 2),
        (320, (209, 152, 48, 90), 2),
        (230, (198, 83, 45, 105), 3),
        (140, (49, 107, 105, 120), 2),
    ]:
        draw.ellipse([cx - r, cy - r, cx + r, cy + r], outline=col, width=width)

    # Subtle axis hairlines
    draw.line([(cx - 430, cy), (cx + 430, cy)], fill=(62, 74, 96, 80), width=2)
    draw.line([(cx, cy - 430), (cx, cy + 430)], fill=(62, 74, 96, 80), width=2)

    # Rising story arc curve positioned cleanly in the upper canvas
    pts = []
    for i in range(120):
        t = i / 119.0
        x = int(95 + t * (w - 190))
        y_norm = 0.80 - 0.58 * (t ** 1.25) - 0.12 * math.sin(t * math.pi * 2.5) * (t ** 0.7)
        y = int(y_norm * h)
        pts.append((x, y))

    for i in range(len(pts) - 1):
        draw.line([pts[i], pts[i + 1]], fill=(209, 152, 48, 235), width=6)

    # Key plot nodes on the curve
    milestone_indices = [6, 34, 62, 86, 108]
    labels = [
        ("발단 (Spark)", -50, -48),
        ("갈등 (Conflict)", -70, -48),
        ("전환 (Midpoint)", -40, 22),
        ("위기 (Crisis)", -115, -36),
        ("절정 (Climax)", -110, -48),
    ]
    font_lbl = get_pil_font(25, bold=True)

    for idx, (lbl, ox, oy) in zip(milestone_indices, labels):
        px, py = pts[idx]
        draw.ellipse([px - 18, py - 18, px + 18, py + 18], fill=(198, 83, 45, 95))
        draw.ellipse([px - 9, py - 9, px + 9, py + 9], fill=(248, 245, 240, 255), outline=(198, 83, 45, 255), width=3)
        draw.text((px + ox, py + oy), lbl, font=font_lbl, fill=(238, 233, 224, 240))

    path = os.path.join(ASSET_DIR, "cover_art.png")
    img.save(path)
    return path


def generate_iceberg_diagram():
    """Generates a high-res Character Iceberg diagram for Slide 4."""
    w, h = 1050, 960
    img = Image.new("RGBA", (w, h), (238, 233, 224, 255))
    draw = ImageDraw.Draw(img)

    # Sky / surface background
    draw.rectangle([0, 0, w, h], fill=(244, 240, 233, 255))
    water_y = 300
    draw.rectangle([0, water_y, w, h], fill=(222, 233, 234, 255))

    # Iceberg polygon (above water - wider peak so text breathes)
    top_poly = [
        (w // 2 - 95, 55),
        (w // 2 + 95, 55),
        (w // 2 + 260, water_y),
        (w // 2 - 260, water_y),
    ]
    draw.polygon(top_poly, fill=(255, 255, 255, 255))
    draw.line([top_poly[0], top_poly[1], top_poly[2], top_poly[3], top_poly[0]], fill=(198, 83, 45, 255), width=4)

    # Iceberg polygon (below water - 3 generous trapezoidal tiers)
    tier1 = [
        (w // 2 - 260, water_y),
        (w // 2 + 260, water_y),
        (w // 2 + 340, 500),
        (w // 2 - 340, 500),
    ]
    tier2 = [
        (w // 2 - 340, 500),
        (w // 2 + 340, 500),
        (w // 2 + 290, 700),
        (w // 2 - 290, 700),
    ]
    tier3 = [
        (w // 2 - 290, 700),
        (w // 2 + 290, 700),
        (w // 2 + 190, 905),
        (w // 2 - 190, 905),
    ]
    draw.polygon(tier1, fill=(75, 130, 128, 235), outline=(255, 255, 255, 190))
    draw.polygon(tier2, fill=(49, 107, 105, 245), outline=(255, 255, 255, 190))
    draw.polygon(tier3, fill=(26, 33, 46, 255), outline=(255, 255, 255, 190))

    # Waterline dashed line
    for x in range(28, w - 28, 28):
        draw.line([(x, water_y), (x + 16, water_y)], fill=(198, 83, 45, 255), width=4)

    f_badge = get_pil_font(21, bold=True)
    f_title = get_pil_font(27, bold=True)
    f_sub = get_pil_font(20, bold=False)

    # Labels on left margin
    draw.text((24, 24), "▲ 수면 위 (가시적 영역 20%)", font=f_badge, fill=(198, 83, 45, 255))
    draw.text((24, water_y + 16), "▼ 수면 아래 (내면적 동력 80%)", font=f_badge, fill=(49, 107, 105, 255))

    def center_text(y_pos, text, font, color):
        bbox = draw.textbbox((0, 0), text, font=font)
        tw = bbox[2] - bbox[0]
        draw.text(((w - tw) // 2, y_pos), text, font=font, fill=color)

    center_text(115, "외면적 특징과 행동", f_title, (26, 33, 46, 255))
    center_text(162, "직업 · 말투 · 습관 · 가시적 사건", f_sub, (85, 94, 108, 255))
    center_text(215, "외면적 목표 (Want)", f_badge, (198, 83, 45, 255))

    center_text(360, "내면의 결핍과 두려움 (Need)", f_title, (255, 255, 255, 255))
    center_text(408, "인물이 진짜 회복해야 할 내면적 가치", f_sub, (235, 242, 242, 255))

    center_text(558, "거짓 믿음 (The Lie)", f_title, (255, 255, 255, 255))
    center_text(606, "세상과 자신을 왜곡해서 보는 방어기제", f_sub, (220, 232, 232, 255))

    center_text(755, "과거의 상처 (The Ghost)", f_title, (209, 152, 48, 255))
    center_text(805, "거짓 믿음을 만든 근원적 트라우마", f_sub, (238, 233, 224, 255))

    path = os.path.join(ASSET_DIR, "character_iceberg.png")
    img.save(path)
    return path


def generate_plot_arc_chart():
    """Generates a high-resolution 3-Act Narrative Tension Curve diagram for Slide 5."""
    w, h = 2200, 600
    img = Image.new("RGBA", (w, h), (255, 255, 255, 255))
    draw = ImageDraw.Draw(img)

    m_left, m_right, m_top, m_bot = 115, 85, 58, 76
    pw = w - m_left - m_right
    ph = h - m_top - m_bot

    act1_x = m_left + int(pw * 0.25)
    act2_x = m_left + int(pw * 0.75)
    mid_x = m_left + int(pw * 0.50)

    # Subtle background fills for Act 1, Act 2, Act 3
    draw.rectangle([m_left, m_top, act1_x, m_top + ph], fill=(248, 245, 240, 255))
    draw.rectangle([act1_x, m_top, act2_x, m_top + ph], fill=(250, 236, 230, 140))
    draw.rectangle([act2_x, m_top, m_left + pw, m_top + ph], fill=(232, 242, 241, 180))

    f_axis = get_pil_font(21, bold=False)
    f_act = get_pil_font(24, bold=True)
    for frac, label in [(0.12, "극한 긴장"), (0.47, "중간 긴장"), (0.82, "일상 균형")]:
        gy = m_top + int(ph * frac)
        draw.line([(m_left, gy), (m_left + pw, gy)], fill=(218, 211, 199, 255), width=2)
        draw.text((16, gy - 14), label, font=f_axis, fill=(135, 142, 153, 255))

    # Vertical dividers for Acts
    for vx in [act1_x, mid_x, act2_x]:
        for y in range(m_top, m_top + ph, 18):
            draw.line([(vx, y), (vx, min(y + 9, m_top + ph))], fill=(180, 172, 160, 255), width=3)

    # Act Headers
    draw.text((m_left + 26, m_top + 12), "1막: 설정과 출발 (0% ~ 25%)", font=f_act, fill=(85, 94, 108, 255))
    draw.text((act1_x + 26, m_top + 12), "2막: 대립·역전·위기 (25% ~ 75%)", font=f_act, fill=(198, 83, 45, 255))
    draw.text((act2_x + 26, m_top + 12), "3막: 절정과 해소 (75% ~ 100%)", font=f_act, fill=(49, 107, 105, 255))

    # Control points: (x_frac, y_frac, title, pct, label_offset_x, label_offset_y)
    key_pts = [
        (0.02, 0.82, "일상 세계", "0%", 10, -52),
        (0.12, 0.62, "① 촉발 사건\n(Inciting Incident)", "12%", 0, -78),
        (0.25, 0.46, "② 1차 관문 통과\n(Plot Point 1)", "25%", 0, -78),
        (0.50, 0.28, "③ 중간점 반전\n(Midpoint)", "50%", 0, -78),
        (0.75, 0.72, "④ 암흑의 순간\n(All Hope Is Lost)", "75%", 0, 24),
        (0.89, 0.14, "⑤ 클라이맥스\n(Climax)", "90%", -125, -10),
        (0.98, 0.56, "새로운 균형\n(Resolution)", "100%", -65, 24),
    ]

    pixel_pts = [(m_left + int(px * pw), m_top + int(py * ph)) for px, py, _, _, _, _ in key_pts]

    curve = []
    for i in range(len(pixel_pts) - 1):
        x0, y0 = pixel_pts[i]
        x1, y1 = pixel_pts[i + 1]
        steps = 45
        for s in range(steps):
            t = s / float(steps)
            t_smooth = (1 - math.cos(t * math.pi)) / 2.0
            cx = int(x0 + (x1 - x0) * t)
            cy = int(y0 + (y1 - y0) * t_smooth)
            curve.append((cx, cy))
    curve.append(pixel_pts[-1])

    for i in range(len(curve) - 1):
        draw.line([curve[i], curve[i + 1]], fill=(198, 83, 45, 255), width=8)

    f_node = get_pil_font(22, bold=True)
    f_pct = get_pil_font(20, bold=True)
    for idx, ((px, py), (_, _, title, pct, ox, oy)) in enumerate(zip(pixel_pts, key_pts)):
        is_major = idx in (1, 2, 3, 4, 5)
        r = 15 if is_major else 10
        col = (198, 83, 45, 255) if is_major else (26, 33, 46, 255)
        draw.ellipse([px - r - 5, py - r - 5, px + r + 5, py + r + 5], fill=(255, 255, 255, 235))
        draw.ellipse([px - r, py - r, px + r, py + r], fill=col)

        lines = title.split("\n")
        for li, line in enumerate(lines):
            bbox = draw.textbbox((0, 0), line, font=f_node)
            lw = bbox[2] - bbox[0]
            draw.text((px - lw // 2 + ox, py + oy + li * 28), line, font=f_node, fill=(26, 33, 46, 255))

        draw.text((px - 22, m_top + ph + 16), pct, font=f_pct, fill=(85, 94, 108, 255))

    draw.line([(m_left, m_top + ph), (m_left + pw, m_top + ph)], fill=(26, 33, 46, 255), width=3)

    path = os.path.join(ASSET_DIR, "plot_arc_chart.png")
    img.save(path)
    return path


# ---------------------------------------------------------------------------
# PPTX Helper Functions
# ---------------------------------------------------------------------------
def add_textbox(slide, left, top, width, height, text="", size_pt=12, bold=False, color=INK_NAVY,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.25, margin_in=0.06):
    """ppt_module.textbox.add_textbox 로 위임하는 호환 래퍼.

    호출부 좌표/인자 순서는 그대로 두고 구현만 ppt_module로 교체했다.
    이로써 두 가지 레이아웃 버그가 동시에 고쳐진다:
    1. text 안의 "\\n"이 이제 진짜 <a:br/> 줄바꿈으로 렌더링된다 --
       기존에는 하나의 run에 리터럴 개행문자로 박혀 PowerPoint/LibreOffice가
       줄바꿈 없이 이어붙여 렌더링했다 (겹침/줄 수 어긋남의 주 원인).
    2. 지정한 height보다 실제 줄바꿈된 텍스트가 더 크면, 조용히 박스
       밖으로 넘치는 대신 글자 크기를 자동으로 줄이고 콘솔에 경고를 남긴다.
    """
    result = _module_add_textbox(
        slide, left, top, width, text, size_pt,
        height=height, bold=bold, color=color, align=align, anchor=anchor,
        line_spacing=line_spacing, margin_in=margin_in, font_name=FONT_KR,
    )
    return result.shape, result.text_frame


def add_header_and_footer(slide, section_kicker, action_title, subtitle, page_num, total_pages=10, dark=False):
    """Consistent editorial header and footer across content slides.

    Kicker/title/subtitle used to sit at hand-picked fixed y-offsets that
    assumed a single short line each; a longer kicker/title on some slides
    overflowed its box height and visually collided with the line below it.
    Heights are now measured from the real text via text_metrics, and each
    box is placed right after the previous one's *actual* bottom edge, so
    this can no longer happen regardless of how long the strings are.
    """
    bg_col = BG_DARK if dark else BG_CREAM
    title_col = INK_LIGHT if dark else INK_NAVY
    sub_col = INK_LIGHT_MUTED if dark else INK_SLATE
    kicker_col = ACCENT_GOLD if dark else ACCENT_RUST
    line_col = BORDER_DARK if dark else BORDER_WARM

    set_slide_background(slide, bg_col)

    header_left, header_w, header_margin, gap = 0.84, 11.8, 0.02, 0.09
    inner_w = header_w - 2 * header_margin
    y = 0.36

    kicker_h = max(measure_block_height_in(section_kicker, 10.5, inner_w, bold=True) + 2 * header_margin * 0.7, 0.24)
    add_textbox(slide, header_left, y, header_w, kicker_h, text=section_kicker,
                size_pt=10.5, bold=True, color=kicker_col, margin_in=header_margin)
    y += kicker_h + gap

    title_h = max(measure_block_height_in(action_title, 21.5, inner_w, bold=True) + 2 * header_margin * 0.7, 0.32)
    add_textbox(slide, header_left, y, header_w, title_h, text=action_title,
                size_pt=21.5, bold=True, color=title_col, margin_in=header_margin)
    y += title_h + gap

    subtitle_h = max(measure_block_height_in(subtitle, 12.0, inner_w) + 2 * header_margin * 0.7, 0.26)
    add_textbox(slide, header_left, y, header_w, subtitle_h, text=subtitle,
                size_pt=12.0, bold=False, color=sub_col, margin_in=header_margin)
    y += subtitle_h

    # Top accent bar spans the measured header block instead of a fixed guess
    add_rect(slide, 0.65, 0.42, 0.08, max(y - 0.42 - gap, 0.6), fill_color=kicker_col)

    # Thin divider rule under header
    divider_y = y + 0.08
    add_rect(slide, 0.65, divider_y, 12.03, 0.015, fill_color=line_col)

    # Footer
    add_rect(slide, 0.65, 7.05, 12.03, 0.012, fill_color=line_col)
    add_textbox(slide, 0.65, 7.10, 6.0, 0.28,
                text="THE CRAFT OF FICTION WRITING  |  소설 창작 실전 가이드",
                size_pt=9.0, bold=False, color=INK_MUTED, margin_in=0.01)
    add_textbox(slide, 10.68, 7.10, 2.0, 0.28,
                text=f"{page_num:02d} / {total_pages:02d}",
                size_pt=9.5, bold=True, color=kicker_col, align=PP_ALIGN.RIGHT, margin_in=0.01)


# ---------------------------------------------------------------------------
# Slide Builders
# ---------------------------------------------------------------------------
def build_slide_1_cover(prs, cover_art_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, BG_DARK)

    # Left vertical terracotta accent bar
    add_rect(slide, 0.0, 0.0, 0.18, 7.5, fill_color=ACCENT_RUST)

    # Right illustration panel positioned cleanly above the bottom module cards
    slide.shapes.add_picture(cover_art_path, Inches(6.85), Inches(0.28), width=Inches(6.05))

    # Category pill
    add_rect(slide, 0.85, 1.05, 3.35, 0.38, fill_color=BG_DARK_CARD, line_color=ACCENT_GOLD, line_pt=1.2, rounded=True, radius=0.45)
    add_textbox(slide, 0.85, 1.05, 3.35, 0.38, text="CREATIVE WRITING MASTERCLASS",
                size_pt=10.5, bold=True, color=ACCENT_GOLD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Main Title -- 제목 높이를 실제 텍스트(2줄 포함)로 측정해 아래 밑줄·
    # 부제를 그 뒤에 순차 배치한다. 고정 높이 1.85in은 42pt 두 줄에는
    # 부족해 부제와 겹치는 원인이었다.
    title_text = "매혹적인\n소설 쓰기의 기술"
    title_h = measure_block_height_in(title_text, 42, 6.2 - 0.04, bold=True, line_spacing=1.18) + 0.03
    add_textbox(slide, 0.85, 1.68, 6.2, title_h,
                text=title_text,
                size_pt=42, bold=True, color=INK_LIGHT, line_spacing=1.18, margin_in=0.02)
    cover_y = 1.68 + title_h

    # Accent underline
    cover_y += 0.10
    add_rect(slide, 0.85, cover_y, 1.6, 0.07, fill_color=ACCENT_RUST)
    cover_y += 0.07 + 0.10

    # Subtitle
    subtitle_text = "아이디어의 첫 불꽃부터 마지막 문장의 퇴고까지,\n독자의 밤을 빼앗는 서사 설계의 10가지 핵심 원칙"
    subtitle_h = measure_block_height_in(subtitle_text, 15, 6.0 - 0.04, line_spacing=1.4) + 0.03
    add_textbox(slide, 0.85, cover_y, 6.0, subtitle_h,
                text=subtitle_text,
                size_pt=15, bold=False, color=INK_LIGHT_MUTED, line_spacing=1.4, margin_in=0.02)

    # Bottom 4 module preview cards
    modules = [
        ("PART 01", "발상과 캐릭터", "로그라인 · 욕망과 결핍"),
        ("PART 02", "플롯과 갈등", "3막 구조 · 씬의 역학"),
        ("PART 03", "시점과 문체", "POV · Show, Don't Tell"),
        ("PART 04", "퇴고와 완성", "3단계 수정 · 작가 루틴"),
    ]
    for i, (part, title, desc) in enumerate(modules):
        cx = 0.85 + i * 2.95
        cy = 5.45
        add_rect(slide, cx, cy, 2.75, 1.38, fill_color=BG_DARK_CARD, line_color=BORDER_DARK, line_pt=1.0, rounded=True, radius=0.06)
        add_rect(slide, cx, cy, 2.75, 0.06, fill_color=ACCENT_GOLD if i % 2 == 1 else ACCENT_RUST)
        add_textbox(slide, cx + 0.16, cy + 0.14, 2.4, 0.26, text=part, size_pt=9.5, bold=True, color=ACCENT_GOLD, margin_in=0.01)
        add_textbox(slide, cx + 0.16, cy + 0.44, 2.4, 0.36, text=title, size_pt=14.5, bold=True, color=INK_LIGHT, margin_in=0.01)
        add_textbox(slide, cx + 0.16, cy + 0.86, 2.4, 0.38, text=desc, size_pt=10.5, bold=False, color=INK_LIGHT_MUTED, margin_in=0.01)


def build_slide_2_pillars(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "01  |  FICTION ARCHITECTURE",
        "완성도 높은 소설은 소재·인물·플롯·문체의 4대 축으로 작동한다",
        "네 가지 기둥 중 하나라도 무너지면 서사의 설득력이 급격히 약화된다",
        page_num=2,
    )

    pillars = [
        (
            "01",
            "매혹적인 전제 (Premise)",
            "호기심을 점화하는 '만약에(What if?)'의 질문",
            "• 익숙한 일상 속 낯선 균열을 포착해 독자의 질문을 유도\n"
            "• 단순한 설정 나열이 아닌 주인공이 겪을 윤리적·현실적 딜레마 내포\n"
            "• 한 문장의 로그라인으로 압축될 때 서사의 방향성이 흔들리지 않음",
            "핵심 질문: 독자가 첫 페이지를 펼칠 이유는 무엇인가?",
            ACCENT_RUST,
            SOFT_RUST_BG,
        ),
        (
            "02",
            "입체적 인물 (Character)",
            "욕망(Want)과 결핍(Need)이 충돌하는 살아있는 주체",
            "• 완벽한 영웅보다 결함과 모순을 지닌 인물에게 독자가 감정이입\n"
            "• 외면적 목표를 추구하는 과정에서 내면의 거짓 믿음이 깨지는 성장\n"
            "• 주인공의 약점을 정확히 파고드는 대립자(Antagonist)와의 상호작용",
            "핵심 질문: 이 인물은 무엇을 간절히 원하며 무엇을 잃을까 두려워하는가?",
            ACCENT_TEAL,
            SOFT_TEAL_BG,
        ),
        (
            "03",
            "인과적 플롯 (Plot & Conflict)",
            "'그리고'가 아닌 '그래서/그러나'로 연결된 사건의 사슬",
            "• 우연에 기댄 해결(데우스 엑스 마키나)을 배제하고 인물의 선택으로 전개\n"
            "• 점진적으로 판돈(Stakes)을 높여 주인공을 퇴로 없는 막다른 골목으로 압박\n"
            "• 3막 구조의 5대 전환점을 통해 긴장과 이완의 서사적 리듬을 설계",
            "핵심 질문: 이전 사건이 다음 사건을 필연적으로 강제하는가?",
            ACCENT_GOLD,
            SOFT_GOLD_BG,
        ),
        (
            "04",
            "감각적 장면과 문체 (Scene & Voice)",
            "요약 설명(Tell)을 넘어 현장을 체험하게 하는 묘사(Show)",
            "• 오감의 디테일과 구체적 행동으로 인물의 감정선을 시각화\n"
            "• 인물의 성격과 관계의 권력 구도가 묻어나는 입체적인 대사 설계\n"
            "• 불필요한 수식어를 덜어내고 장면의 속도에 맞춘 문장 호흡 조율",
            "핵심 질문: 독자의 머릿속에 이 장면이 영화처럼 그려지는가?",
            ACCENT_BURGUNDY,
            SOFT_RUST_BG,
        ),
    ]

    positions = [
        (0.65, 1.60),
        (6.78, 1.60),
        (0.65, 4.02),
        (6.78, 4.02),
    ]

    for (num, title, sub, body, key_q, col, tint), (x, y) in zip(pillars, positions):
        w, h = 5.90, 2.24
        add_rect(slide, x, y, w, h, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.04)
        add_rect(slide, x, y, 0.10, h, fill_color=col)

        add_rect(slide, x + 0.26, y + 0.18, 0.56, 0.44, fill_color=col, rounded=True, radius=0.15)
        add_textbox(slide, x + 0.26, y + 0.18, 0.56, 0.44, text=num,
                    size_pt=14, bold=True, color=BG_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        add_textbox(slide, x + 0.94, y + 0.14, 4.8, 0.34, text=title,
                    size_pt=15.5, bold=True, color=INK_NAVY, margin_in=0.01)
        add_textbox(slide, x + 0.94, y + 0.49, 4.8, 0.26, text=sub,
                    size_pt=10.8, bold=True, color=col, margin_in=0.01)

        add_textbox(slide, x + 0.26, y + 0.78, 5.45, 0.95, text=body,
                    size_pt=10.8, bold=False, color=INK_SLATE, line_spacing=1.28, margin_in=0.02)

        add_rect(slide, x + 0.22, y + 1.76, 5.50, 0.36, fill_color=tint, rounded=True, radius=0.10)
        add_textbox(slide, x + 0.32, y + 1.77, 5.30, 0.34, text=key_q,
                    size_pt=10.2, bold=True, color=INK_NAVY, anchor=MSO_ANCHOR.MIDDLE, margin_in=0.01)

    add_rect(slide, 0.65, 6.42, 12.03, 0.48, fill_color=INK_NAVY, rounded=True, radius=0.08)
    add_textbox(slide, 0.85, 6.42, 11.6, 0.48,
                text="핵심 공식:  매력적인 소재(Premise)가 독자를 부르고, 입체적인 인물(Character)이 페이지를 넘기게 하며, 단단한 플롯과 문체가 여운을 완성한다.",
                size_pt=11.5, bold=True, color=INK_LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def build_slide_3_logline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "02  |  IDEA & LOGLINE DESIGN",
        "막연한 소재를 한 문장의 로그라인(Logline)으로 압축해 서사의 나침반을 세운다",
        "좋은 로그라인은 주인공, 촉발 사건, 핵심 목표, 그리고 실패 시 치러야 할 대가(Stakes)를 명확히 보여준다",
        page_num=3,
    )

    add_rect(slide, 0.65, 1.56, 12.03, 1.52, fill_color=BG_SAND, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.04)
    add_textbox(slide, 0.88, 1.64, 4.0, 0.28, text="강력한 로그라인의 4요소 공식", size_pt=11.5, bold=True, color=ACCENT_RUST, margin_in=0.01)

    formula_steps = [
        ("1. 결핍을 지닌 주인공", "어떤 약점·상황에 처한 인물이", BG_WHITE, INK_NAVY),
        ("2. 촉발 사건 (Inciting)", "예기치 못한 사건을 겪은 뒤", BG_WHITE, INK_NAVY),
        ("3. 핵심 목표와 대립", "강력한 적수/장애물에 맞서", BG_WHITE, INK_NAVY),
        ("4. 치명적 대가 (Stakes)", "무엇을 걸고 싸우는가?", SOFT_RUST_BG, ACCENT_RUST),
    ]
    for i, (ftitle, fdesc, fbg, fcol) in enumerate(formula_steps):
        bx = 0.88 + i * 2.94
        by = 1.98
        add_rect(slide, bx, by, 2.70, 0.92, fill_color=fbg, line_color=ACCENT_RUST if i == 3 else BORDER_WARM, line_pt=1.2, rounded=True, radius=0.08)
        add_textbox(slide, bx + 0.14, by + 0.12, 2.42, 0.32, text=ftitle, size_pt=12.0, bold=True, color=fcol, margin_in=0.01)
        add_textbox(slide, bx + 0.14, by + 0.46, 2.42, 0.34, text=fdesc, size_pt=10.5, bold=False, color=INK_SLATE, margin_in=0.01)
        if i < 3:
            add_textbox(slide, bx + 2.70, by + 0.24, 0.24, 0.40, text="+", size_pt=18, bold=True, color=ACCENT_RUST, align=PP_ALIGN.CENTER)

    add_textbox(slide, 0.65, 3.24, 12.0, 0.32,
                text="실전 변환 예시: 추상적 소재에서 구체적 로그라인으로",
                size_pt=13.5, bold=True, color=INK_NAVY, margin_in=0.01)

    examples = [
        (
            "장르 01 · 미스터리 스릴러",
            "기억상실증에 걸린 형사가 살인사건을 수사하는 이야기",
            "단기 기억이 24시간마다 초기화되는 강력계 형사가, 자신의 수첩에 적힌 마지막 용의자가 바로 자신임을 알게 되면서 기억이 사라지기 전에 진범을 찾아야 하는 사투",
            "제한 시간(24시간) + 아이러니(자신이 용의자)",
        ),
        (
            "장르 02 · SF / 판타지",
            "감정을 사고파는 미래 사회에서 벌어지는 청춘 드라마",
            "타인의 슬픔을 대신 겪어주고 돈을 버는 ‘감정 대리인’ 청년이, 의뢰인의 지워진 기억 속에서 10년 전 실종된 여동생의 흔적을 발견하고 거대 기업의 금기를 추적하는 이야기",
            "독창적 직업 + 개인적 절박함(실종된 동생)",
        ),
        (
            "장르 03 · 휴먼 / 가족 드라마",
            "사이가 나쁜 모녀가 함께 여행을 떠나며 화해하는 소설",
            "평생 연락을 끊고 살던 어머니의 유언장 조건 때문에, 빚더미에 오른 딸이 어머니의 첫사랑을 찾아 일주일간 동행하며 몰랐던 어머니의 청춘과 마주하는 여정",
            "강제적 동행(유언 조건) + 내면적 화해",
        ),
    ]

    for idx, (genre, weak, strong, point) in enumerate(examples):
        ry = 3.64 + idx * 1.08
        add_rect(slide, 0.65, ry, 4.15, 0.96, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.05)
        add_textbox(slide, 0.80, ry + 0.08, 3.85, 0.24, text=f"[막연한 소재]  {genre}", size_pt=9.8, bold=True, color=INK_MUTED, margin_in=0.01)
        add_textbox(slide, 0.80, ry + 0.36, 3.85, 0.50, text=f'“{weak}”', size_pt=11.0, bold=False, color=INK_SLATE, margin_in=0.01)

        add_textbox(slide, 4.82, ry + 0.26, 0.42, 0.44, text="→", size_pt=20, bold=True, color=ACCENT_RUST, align=PP_ALIGN.CENTER)

        add_rect(slide, 5.26, ry, 7.42, 0.96, fill_color=BG_WHITE, line_color=ACCENT_RUST, line_pt=1.3, rounded=True, radius=0.05)
        add_rect(slide, 5.26, ry, 0.09, 0.96, fill_color=ACCENT_RUST)
        add_textbox(slide, 5.46, ry + 0.08, 4.2, 0.24, text="[완성된 로그라인]", size_pt=9.8, bold=True, color=ACCENT_RUST, margin_in=0.01)
        add_rect(slide, 9.75, ry + 0.08, 2.78, 0.27, fill_color=SOFT_RUST_BG, rounded=True, radius=0.2)
        add_textbox(slide, 9.75, ry + 0.08, 2.78, 0.27, text=point, size_pt=9.0, bold=True, color=ACCENT_RUST, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
        add_textbox(slide, 5.46, ry + 0.36, 7.05, 0.54, text=strong, size_pt=10.6, bold=True, color=INK_NAVY, line_spacing=1.22, margin_in=0.01)


def build_slide_4_character(prs, iceberg_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "03  |  CHARACTER PSYCHOLOGY",
        "독자가 사랑하는 인물은 외면적 목표(Want)와 내면적 결핍(Need)의 간극에서 탄생한다",
        "인물의 프로필을 나열하는 대신, 과거의 상처(Ghost)가 만든 거짓 믿음(Lie)과 행동 동기를 설계하라",
        page_num=4,
    )

    add_rect(slide, 0.65, 1.56, 5.35, 5.32, fill_color=BG_SAND, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.03)
    add_textbox(slide, 0.85, 1.65, 4.95, 0.32, text="인물 설계의 빙산 모델 (Character Iceberg)", size_pt=13.0, bold=True, color=INK_NAVY, margin_in=0.01)
    slide.shapes.add_picture(iceberg_path, Inches(0.80), Inches(2.02), width=Inches(5.05))

    right_cards = [
        (
            "욕망(Want) vs 필요(Need)의 충돌",
            "주인공은 눈에 보이는 목표(성공·복수·탈출)를 향해 달리지만, 결말에 이르러 자신이 진짜 필요했던 것(신뢰·용서·진실)을 깨닫고 변화한다.",
            ACCENT_RUST,
        ),
        (
            "적대자(Antagonist)는 주인공의 거울",
            "최고의 악역은 스스로를 선이라고 믿는 인물이다. 주인공과 같은 문제에 직면했지만 정반대의 선택을 한 그림자로 설계할 때 대립이 깊어진다.",
            ACCENT_TEAL,
        ),
        (
            "성격은 형용사가 아니라 ‘선택’으로 증명",
            "평화로운 순간의 친절함은 성격을 보여주지 못한다. 가장 소중한 두 가치 중 하나를 포기해야 하는 극한의 딜레마에서 내리는 선택이 진짜 캐릭터다.",
            ACCENT_GOLD,
        ),
    ]

    for i, (ctitle, cbody, ccol) in enumerate(right_cards):
        cy = 1.56 + i * 1.12
        add_rect(slide, 6.22, cy, 6.46, 1.00, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.04)
        add_rect(slide, 6.22, cy, 0.09, 1.00, fill_color=ccol)
        add_textbox(slide, 6.44, cy + 0.10, 6.1, 0.28, text=ctitle, size_pt=13.0, bold=True, color=INK_NAVY, margin_in=0.01)
        add_textbox(slide, 6.44, cy + 0.40, 6.1, 0.54, text=cbody, size_pt=10.6, bold=False, color=INK_SLATE, line_spacing=1.24, margin_in=0.01)

    add_rect(slide, 6.22, 4.96, 6.46, 1.92, fill_color=INK_NAVY, rounded=True, radius=0.04)
    add_textbox(slide, 6.44, 5.06, 6.0, 0.28,
                text="실전 캐릭터 설계 시트 예시  —  <폐기처분 복원가 ‘한서진’>",
                size_pt=12.0, bold=True, color=ACCENT_GOLD, margin_in=0.01)

    sheet_items = [
        ("과거의 상처 (Ghost)", "어린 시절 자신의 거짓말로 가족의 유품을 잃어버린 죄책감"),
        ("거짓 믿음 (The Lie)", "“사람은 결국 배신하지만, 사물은 절대 거짓말을 하지 않는다.”"),
        ("외면적 목표 (Want)", "불타버린 고택의 미공개 유작을 완벽히 복원해 거액의 보수를 받는 것"),
        ("내면적 성장 (Need)", "완벽한 복원보다 남겨진 유가족의 상처를 보듬고 진실을 마주하는 용기"),
    ]
    for idx, (k, v) in enumerate(sheet_items):
        sy = 5.42 + idx * 0.34
        add_textbox(slide, 6.44, sy, 1.85, 0.30, text=f"• {k}", size_pt=10.0, bold=True, color=INK_LIGHT, margin_in=0.01)
        add_textbox(slide, 8.32, sy, 4.25, 0.30, text=v, size_pt=10.0, bold=False, color=INK_LIGHT_MUTED, margin_in=0.01)


def build_slide_5_plot_arc(prs, plot_chart_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "04  |  THREE-ACT STRUCTURE & PLOT BEATS",
        "3막 구조와 5대 전환점(Plot Beats)으로 독자가 책을 덮지 못하는 긴장감을 조율한다",
        "서사의 황금 비율(1막 25% : 2막 50% : 3막 25%)을 통해 위기의 고조와 해소의 쾌감을 정교하게 배치한다",
        page_num=5,
    )

    add_rect(slide, 0.65, 1.54, 12.03, 3.36, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.03)
    slide.shapes.add_picture(plot_chart_path, Inches(0.75), Inches(1.60), width=Inches(11.83))

    beats = [
        ("① 촉발 사건 (12%)", "일상의 균열", "주인공의 평온한 일상을 깨뜨리고 이야기 속으로 호출하는 사건", ACCENT_RUST),
        ("② 1차 관문 (25%)", "돌아갈 다리 불태우기", "주인공이 기존 세계를 떠나 본격적인 모험·대립의 세계로 진입하는 결단", ACCENT_TEAL),
        ("③ 중간점 반전 (50%)", "시야의 전환", "숨겨진 진실이 드러나며 수동적 대응에서 능동적 공세로 전환되는 지점", ACCENT_GOLD),
        ("④ 암흑의 순간 (75%)", "최악의 패배", "모든 계획이 무너지고 가장 소중한 것을 잃을 위기에 처하는 절망의 바닥", ACCENT_BURGUNDY),
        ("⑤ 클라이맥스 (90%)", "최후의 대결과 각성", "내면의 거짓 믿음을 극복한 주인공이 자신의 힘으로 핵심 갈등을 매듭짓는 순간", INK_NAVY),
    ]

    for i, (btitle, bsub, bdesc, bcol) in enumerate(beats):
        bx = 0.65 + i * 2.44
        by = 5.06
        bw, bh = 2.31, 1.82
        add_rect(slide, bx, by, bw, bh, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.05)
        add_rect(slide, bx, by, bw, 0.08, fill_color=bcol)
        add_textbox(slide, bx + 0.12, by + 0.14, bw - 0.24, 0.28, text=btitle, size_pt=11.5, bold=True, color=bcol, margin_in=0.01)
        add_textbox(slide, bx + 0.12, by + 0.42, bw - 0.24, 0.26, text=bsub, size_pt=11.0, bold=True, color=INK_NAVY, margin_in=0.01)
        add_textbox(slide, bx + 0.12, by + 0.74, bw - 0.24, 0.96, text=bdesc, size_pt=10.0, bold=False, color=INK_SLATE, line_spacing=1.24, margin_in=0.01)


def build_slide_6_scene_dynamics(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "05  |  SCENE & SEQUEL DYNAMICS",
        "모든 장면(Scene)은 목표·충돌·전환으로 전진하고 반추(Sequel)로 호흡을 조절한다",
        "외면적 사건의 속도감(Scene)과 내면적 감정 소화(Sequel)를 교차 배치해 독자의 피로도를 막고 몰입을 깊게 만든다",
        page_num=6,
    )

    add_rect(slide, 0.65, 1.56, 7.45, 5.32, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.03)
    add_textbox(slide, 0.88, 1.68, 7.0, 0.34,
                text="서사 전진의 양대 엔진: 장면(Scene)과 반추(Sequel)의 교차 리듬",
                size_pt=13.5, bold=True, color=INK_NAVY, margin_in=0.01)

    add_rect(slide, 0.88, 2.12, 7.00, 2.15, fill_color=SOFT_RUST_BG, line_color=ACCENT_RUST, line_pt=1.2, rounded=True, radius=0.04)
    add_textbox(slide, 1.06, 2.22, 6.6, 0.30,
                text="ENGINE A  |  외면적 행동 장면 (Action Scene) — 사건을 발생시키고 긴장을 고조",
                size_pt=11.8, bold=True, color=ACCENT_RUST, margin_in=0.01)

    scene_steps = [
        ("1. 목표 (Goal)", "장면 시작 시 주인공이 원하는 구체적이고 가시적인 목적"),
        ("2. 충돌 (Conflict)", "목표 달성을 가로막는 상대방이나 환경과의 치열한 공방"),
        ("3. 전환/재난 (Disaster)", "‘예스, 하지만(Yes, but)’ 또는 ‘노, 심지어(No, and furthermore)’로 끝나는 위기"),
    ]
    for i, (st, sd) in enumerate(scene_steps):
        sx = 1.06 + i * 2.24
        add_rect(slide, sx, 2.60, 2.12, 1.50, fill_color=BG_WHITE, line_color=BORDER_WARM, rounded=True, radius=0.06)
        add_textbox(slide, sx + 0.12, 2.70, 1.88, 0.30, text=st, size_pt=11.2, bold=True, color=INK_NAVY, margin_in=0.01)
        add_textbox(slide, sx + 0.12, 3.04, 1.88, 0.96, text=sd, size_pt=9.8, bold=False, color=INK_SLATE, line_spacing=1.22, margin_in=0.01)

    add_rect(slide, 0.88, 4.46, 7.00, 2.15, fill_color=SOFT_TEAL_BG, line_color=ACCENT_TEAL, line_pt=1.2, rounded=True, radius=0.04)
    add_textbox(slide, 1.06, 4.56, 6.6, 0.30,
                text="ENGINE B  |  내면적 반추 장면 (Reaction Sequel) — 감정을 납득시키고 다음 목표를 수립",
                size_pt=11.8, bold=True, color=ACCENT_TEAL, margin_in=0.01)

    sequel_steps = [
        ("1. 감정 반응 (Reaction)", "직전 장면의 재난/충격에 대한 인물의 정서적·신체적 여파"),
        ("2. 딜레마 (Dilemma)", "나쁜 선택지들 사이에서 최악을 피하기 위해 고뇌하는 분석 과정"),
        ("3. 새로운 결단 (Decision)", "위험을 감수하고 다음 행동 장면(Scene)의 목표를 세우는 결심"),
    ]
    for i, (st, sd) in enumerate(sequel_steps):
        sx = 1.06 + i * 2.24
        add_rect(slide, sx, 4.94, 2.12, 1.50, fill_color=BG_WHITE, line_color=BORDER_WARM, rounded=True, radius=0.06)
        add_textbox(slide, sx + 0.12, 5.04, 1.88, 0.30, text=st, size_pt=11.2, bold=True, color=INK_NAVY, margin_in=0.01)
        add_textbox(slide, sx + 0.12, 5.38, 1.88, 0.96, text=sd, size_pt=9.8, bold=False, color=INK_SLATE, line_spacing=1.22, margin_in=0.01)

    add_rect(slide, 8.32, 1.56, 4.36, 3.45, fill_color=BG_SAND, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.03)
    add_textbox(slide, 8.54, 1.68, 3.95, 0.32, text="입체적 서사를 만드는 갈등의 3중 레이어", size_pt=13.0, bold=True, color=INK_NAVY, margin_in=0.01)

    layers = [
        ("내적 갈등 (Inner)", "양심 vs 욕망, 두려움 vs 의무 등 인물 내면의 심리적 분열", ACCENT_RUST),
        ("대인 갈등 (Personal)", "연인·가족·동료·숙적과의 이해관계 충돌과 배신", ACCENT_TEAL),
        ("환경/사회 갈등 (Extra)", "재난, 부조리한 제도, 시간 제한 등 거대한 외부 압력", ACCENT_GOLD),
    ]
    for i, (lt, ld, lcol) in enumerate(layers):
        ly = 2.12 + i * 0.92
        add_rect(slide, 8.54, ly, 3.92, 0.80, fill_color=BG_WHITE, line_color=BORDER_WARM, rounded=True, radius=0.05)
        add_rect(slide, 8.54, ly, 0.08, 0.80, fill_color=lcol)
        add_textbox(slide, 8.72, ly + 0.08, 3.6, 0.24, text=lt, size_pt=11.2, bold=True, color=lcol, margin_in=0.01)
        add_textbox(slide, 8.72, ly + 0.34, 3.6, 0.40, text=ld, size_pt=9.8, bold=False, color=INK_SLATE, margin_in=0.01)

    add_rect(slide, 8.32, 5.18, 4.36, 1.70, fill_color=INK_NAVY, rounded=True, radius=0.04)
    add_textbox(slide, 8.54, 5.30, 3.92, 0.28, text="장면 생존 테스트 (Scene Value Test)", size_pt=12.0, bold=True, color=ACCENT_GOLD, margin_in=0.01)
    add_textbox(slide, 8.54, 5.64, 3.92, 1.12,
                text="• 장면의 시작과 끝 사이에 인물의 상황·감정의 부호(+/-)가 바뀌었는가?\n"
                     "• 플롯을 전진시키거나 인물의 핵심 본성을 드러내는가?\n"
                     "→ 둘 다 아니라면 아무리 문장이 아름다워도 과감히 삭제하라.",
                size_pt=10.0, bold=False, color=INK_LIGHT, line_spacing=1.28, margin_in=0.01)


def build_slide_7_pov(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "06  |  POINT OF VIEW (POV) & NARRATIVE DISTANCE",
        "누구의 눈으로 어디까지 보여줄 것인가: 시점(POV)이 소설의 몰입도를 결정한다",
        "각 시점의 정보 통제 범위와 심리적 거리를 이해하고, 한 장면 안에서는 단일 시점을 엄격히 유지하라",
        page_num=7,
    )

    headers = ["시점 유형", "정보 통제와 서술 범위", "서사적 강점 (Pros)", "주의할 함정 (Cons)", "적합한 장르 및 대표 효과"]
    col_widths = [2.05, 2.75, 2.55, 2.45, 2.23]
    rows_data = [
        (
            "1인칭 주인공 시점\n(‘나’의 서술)",
            "오직 ‘나’가 보고 듣고 추측한 정보만 서술 가능",
            "• 독자와 주인공의 즉각적 밀착\n• 화자의 독특한 목소리·유머 극대화\n• ‘신뢰할 수 없는 화자’ 트릭에 유리",
            "• 주인공이 없는 현장의 사건 서술 불가\n• 지나친 자기연민이나 독백 과잉 위험",
            "청춘 성장소설, 느와르,\n심리 미스터리, 고백록",
        ),
        (
            "3인칭 제한 시점\n(현대 소설의 표준)",
            "초점화된 특정 인물(‘그/그녀’)의 내면과 시선에만 밀착",
            "• 1인칭의 깊은 정서적 몰입감과 3인칭의 객관적 묘사력을 동시에 확보\n• 서사적 거리를 유연하게 줌인/줌아웃",
            "• 초점 인물이 눈치채지 못한 자신의 표정·모습을 직접 묘사하는 실수 주의",
            "현대 문학 전반, 장르소설,\n장편 드라마의 가장 안정적 선택",
        ),
        (
            "3인칭 전지적 시점\n(신적 조망)",
            "모든 인물의 내면, 과거·미래, 보이지 않는 사건까지 통찰",
            "• 방대한 세계관과 시대상을 조망\n• 인물 간의 오해와 극적 아이러니를 독자에게 입체적으로 제시",
            "• 한 장면에서 여러 인물의 속마음을 오가면(Head-hopping) 긴장감 급감",
            "대하 역사소설, 가족 사가(Saga),\n우화적·철학적 서사",
        ),
        (
            "다중 교차 시점\n(챕터별 3인칭/1인칭)",
            "챕터나 절(Section) 단위로 시점 인물을 교대하며 서술",
            "• 주인공과 적대자 양측의 정당성 부여\n• 시점 교차 절단신공(Cliffhanger)으로 페이지 터너 효과 극대화",
            "• 시점 인물 간 목소리가 비슷하면 독자가 화자를 혼동\n• 서사 밀도가 낮은 시점 챕터가 지루해짐",
            "군상극, 범죄 스릴러,\n거대 판타지, 로맨스(쌍방 시점)",
        ),
    ]

    tx = 0.65
    ty = 1.56
    cx = tx
    for h_text, cw in zip(headers, col_widths):
        add_rect(slide, cx, ty, cw, 0.46, fill_color=INK_NAVY, line_color=BORDER_WARM, line_pt=0.8)
        add_textbox(slide, cx + 0.08, ty + 0.04, cw - 0.16, 0.38, text=h_text,
                    size_pt=10.8, bold=True, color=INK_LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin_in=0.01)
        cx += cw

    for r_idx, row in enumerate(rows_data):
        ry = ty + 0.46 + r_idx * 1.06
        row_bg = SOFT_RUST_BG if r_idx == 1 else (BG_WHITE if r_idx % 2 == 0 else BG_SAND)
        cx = tx
        for c_idx, (cell_text, cw) in enumerate(zip(row, col_widths)):
            add_rect(slide, cx, ry, cw, 1.06, fill_color=row_bg, line_color=BORDER_WARM, line_pt=0.8)
            is_first = c_idx == 0
            t_col = (ACCENT_RUST if r_idx == 1 else INK_NAVY) if is_first else INK_SLATE
            add_textbox(slide, cx + 0.10, ry + 0.08, cw - 0.20, 0.90, text=cell_text,
                        size_pt=10.5 if is_first else 9.9, bold=is_first, color=t_col,
                        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.22, margin_in=0.01)
            cx += cw

    add_rect(slide, 0.65, 6.38, 12.03, 0.52, fill_color=SOFT_GOLD_BG, line_color=ACCENT_GOLD, line_pt=1.2, rounded=True, radius=0.06)
    add_textbox(slide, 0.85, 6.38, 11.6, 0.52,
                text="초보 작가의 대표적 실수 ‘시점 널뛰기(Head-Hopping)’ 금지:  한 장면 안에서는 카메라를 한 인물의 어깨 위에만 고정하고, 타인의 속마음은 표정·행동으로만 추론하게 하라.",
                size_pt=10.8, bold=True, color=INK_NAVY, anchor=MSO_ANCHOR.MIDDLE, margin_in=0.01)


def build_slide_8_show_dont_tell(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "07  |  STYLE & DESCRIPTION CLINIC (SHOW, DON'T TELL)",
        "감정을 단어로 요약(Tell)하지 말고, 오감과 행동의 디테일로 독자가 직접 느끼게 하라(Show)",
        "추상적인 감정어(슬펐다, 화났다, 불안했다)를 지우고 그 감정이 만들어내는 신체 반응과 사물의 이미지로 치환한다",
        page_num=8,
    )

    clinics = [
        (
            "CASE 01  |  감정과 심리 묘사",
            "“지은은 면접 결과를 기다리며 너무나 불안하고 초조해서 견딜 수가 없었다.”",
            "신체 반응과 사물의 디테일로 치환 (Show)",
            "“지은은 식어버린 머그잔 손잡이를 만지작거리다, 기어이 무릎 위에 올려둔 냅킨 모서리를 손톱 끝으로 잘게 찢어놓았다.”",
            "시각·촉각 행동으로 불안감 체감",
        ),
        (
            "CASE 02  |  인물의 성격 제시",
            "“민수는 피도 눈물도 없이 냉정하고 계산적인 성격의 소유자였다.”",
            "결정적 순간의 행동과 시선 처리 (Show)",
            "“민수는 오열하는 동생의 어깨 너머로 벽시계의 분침을 먼저 확인한 뒤, 지갑에서 명함 한 장만 꺼내 식탁 끝에 소리 없이 내려놓았다.”",
            "시계 확인 + 명함 한 장으로 냉담함 증명",
        ),
        (
            "CASE 03  |  공간과 분위기 구축",
            "“그곳은 오랫동안 사람이 살지 않아 무척 음산하고 불길한 폐가였다.”",
            "청각·후각·촉각의 다중 감각 결합 (Show)",
            "“문지방을 넘자 젖은 곰팡이 냄새가 훅 끼쳤고, 깨진 유리창 틈으로 스며든 바람에 들뜬 벽지가 뱀 허물처럼 바스락거렸다.”",
            "후각(곰팡이) + 청각(벽지 소리) 입체화",
        ),
    ]

    for idx, (case_tag, t_txt, s_lbl, s_txt, note) in enumerate(clinics):
        ry = 1.54 + idx * 1.22
        add_rect(slide, 0.65, ry, 4.45, 1.10, fill_color=BG_SAND, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.04)
        add_textbox(slide, 0.82, ry + 0.08, 4.1, 0.24, text=f"{case_tag}  —  [요약 설명: Tell]", size_pt=9.5, bold=True, color=INK_MUTED, margin_in=0.01)
        add_textbox(slide, 0.82, ry + 0.38, 4.1, 0.64, text=t_txt, size_pt=10.6, bold=False, color=INK_SLATE, line_spacing=1.22, margin_in=0.01)

        add_textbox(slide, 5.12, ry + 0.32, 0.38, 0.44, text="→", size_pt=20, bold=True, color=ACCENT_RUST, align=PP_ALIGN.CENTER)

        add_rect(slide, 5.52, ry, 7.16, 1.10, fill_color=BG_WHITE, line_color=ACCENT_TEAL, line_pt=1.3, rounded=True, radius=0.04)
        add_rect(slide, 5.52, ry, 0.09, 1.10, fill_color=ACCENT_TEAL)
        add_textbox(slide, 5.72, ry + 0.08, 3.8, 0.24, text=f"[{s_lbl}]", size_pt=9.5, bold=True, color=ACCENT_TEAL, margin_in=0.01)
        add_rect(slide, 9.65, ry + 0.08, 2.88, 0.27, fill_color=SOFT_TEAL_BG, rounded=True, radius=0.2)
        add_textbox(slide, 9.65, ry + 0.08, 2.88, 0.27, text=f"✓ {note}", size_pt=9.0, bold=True, color=ACCENT_TEAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
        add_textbox(slide, 5.72, ry + 0.38, 6.8, 0.64, text=s_txt, size_pt=10.8, bold=True, color=INK_NAVY, line_spacing=1.24, margin_in=0.01)

    add_rect(slide, 0.65, 5.28, 12.03, 1.60, fill_color=INK_NAVY, rounded=True, radius=0.04)
    add_textbox(slide, 0.88, 5.38, 11.5, 0.28,
                text="살아 숨 쉬는 대사(Dialogue)를 쓰는 3가지 실전 원칙",
                size_pt=12.0, bold=True, color=ACCENT_GOLD, margin_in=0.01)

    d_rules = [
        ("01. 정보 브리핑 대사 금지", "독자에게 배경설정을 알려주려고 인물끼리 이미 아는 사실을 말하게 하지 마라. 대사는 정보 전달이 아니라 ‘심리전’이다."),
        ("02. 서브텍스트(행간) 활용", "인물이 속마음을 정직하게 다 말하면 긴장이 사라진다. 정작 하고 싶은 말은 숨기고 엉뚱한 화제로 감정을 우회시켜라."),
        ("03. 지문과 행동 비트 결합", "“~라고 슬프게 말했다” 같은 부사형 전달동사를 줄이고, 대사 사이사이에 시선 회피나 침묵 같은 행동 비트(Action Beat)를 넣어라."),
    ]
    for i, (dt, dd) in enumerate(d_rules):
        dx = 0.88 + i * 3.92
        add_rect(slide, dx, 5.72, 3.72, 1.02, fill_color=BG_DARK_CARD, line_color=BORDER_DARK, rounded=True, radius=0.05)
        add_textbox(slide, dx + 0.14, 5.80, 3.44, 0.26, text=dt, size_pt=10.8, bold=True, color=INK_LIGHT, margin_in=0.01)
        add_textbox(slide, dx + 0.14, 6.08, 3.44, 0.60, text=dd, size_pt=9.5, bold=False, color=INK_LIGHT_MUTED, line_spacing=1.20, margin_in=0.01)


def build_slide_9_revision(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "08  |  DRAFTING & 3-PASS REVISION PROCESS",
        "초고는 뜨거운 가슴으로 거침없이 쓰고, 퇴고는 차가운 머리로 3단계에 걸쳐 조각한다",
        "완벽주의 때문에 첫 장만 무한히 고치는 함정에서 벗어나, 구조·장면·문장의 순서로 해상도를 높여가며 다듬는다",
        page_num=9,
    )

    add_rect(slide, 0.65, 1.56, 4.55, 5.32, fill_color=BG_SAND, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.03)
    add_textbox(slide, 0.88, 1.70, 4.1, 0.32, text="초고 완주의 제1원칙", size_pt=13.5, bold=True, color=ACCENT_RUST, margin_in=0.01)
    add_textbox(slide, 0.88, 2.08, 4.1, 1.15,
                text="“초고를 쓸 때는 내면의 편집자를 방문 밖에 가두어라.”\n\n백지 위에서는 그 어떤 퇴고도 불가능하다. 문장이 거칠고 설정에 구멍이 나더라도 일단 결말의 마침표까지 달려가 전체 윤곽을 확보하는 것이 최우선이다.",
                size_pt=11.0, bold=False, color=INK_NAVY, line_spacing=1.32, margin_in=0.01)

    add_rect(slide, 0.88, 3.42, 4.10, 1.65, fill_color=BG_WHITE, line_color=BORDER_WARM, rounded=True, radius=0.05)
    add_textbox(slide, 1.04, 3.52, 3.8, 0.28, text="소설 완성 에너지 배분 비율", size_pt=11.5, bold=True, color=INK_NAVY, margin_in=0.01)

    bx = 1.04
    by = 3.90
    bw = 3.78
    add_rect(slide, bx, by, bw * 0.20, 0.42, fill_color=ACCENT_GOLD)
    add_textbox(slide, bx, by, bw * 0.20, 0.42, text="20%", size_pt=10.5, bold=True, color=BG_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, bx + bw * 0.20, by, bw * 0.30, 0.42, fill_color=ACCENT_TEAL)
    add_textbox(slide, bx + bw * 0.20, by, bw * 0.30, 0.42, text="30%", size_pt=10.5, bold=True, color=BG_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, bx + bw * 0.50, by, bw * 0.50, 0.42, fill_color=ACCENT_RUST)
    add_textbox(slide, bx + bw * 0.50, by, bw * 0.50, 0.42, text="50% (퇴고)", size_pt=10.5, bold=True, color=BG_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    add_textbox(slide, 1.04, 4.42, 3.8, 0.55,
                text="• 기획·구조 설계 (20%)  • 초고 집필 (30%)\n• 3단계 퇴고 및 윤문 (50%) — 진짜 소설은 퇴고에서 태어난다.",
                size_pt=9.8, bold=False, color=INK_SLATE, line_spacing=1.25, margin_in=0.01)

    add_rect(slide, 0.88, 5.24, 4.10, 1.44, fill_color=INK_NAVY, rounded=True, radius=0.05)
    add_textbox(slide, 1.06, 5.34, 3.74, 0.28, text="필수 과정: 원고 숙성기 (Cooling Off)", size_pt=11.5, bold=True, color=ACCENT_GOLD, margin_in=0.01)
    add_textbox(slide, 1.06, 5.66, 3.74, 0.90,
                text="초고 완성 직후에는 머릿속 설정과 실제 문장을 구별하지 못한다. 최소 며칠~2주간 원고를 서랍에 넣어두었다가 ‘낯선 독자의 눈’으로 다시 읽어라.",
                size_pt=10.0, bold=False, color=INK_LIGHT_MUTED, line_spacing=1.25, margin_in=0.01)

    passes = [
        (
            "PASS 01",
            "1차 구조 퇴고 (Macro / Structural Edit)",
            "나무가 아닌 숲을 보며 서사의 뼈대와 인과관계를 재조립하는 단계",
            "✓ 주인공의 욕망과 행동이 시작부터 클라이맥스까지 일관되게 이어지는가?\n"
            "✓ 중반부가 늘어지지 않는가? (기능이 겹치는 장면과 인물을 과감히 통합·삭제)\n"
            "✓ 복선이 회수되었으며, 결말이 우연이 아닌 주인공의 선택으로 해결되는가?",
            ACCENT_RUST,
        ),
        (
            "PASS 02",
            "2차 장면·리듬 퇴고 (Scene & Pacing Edit)",
            "각 장면의 밀도와 대화의 텐션, 시점의 일관성을 조율하는 단계",
            "✓ 각 장면의 첫 문장이 흥미롭고 마지막 문장이 다음 장면을 궁금하게 만드는가?\n"
            "✓ 한 장면 안에서 시점(POV)이 흔들리거나 설명(Tell)이 과도한 구간은 없는가?\n"
            "✓ 긴박한 장면에서는 단문으로 속도를 높이고, 정서적 장면에서는 호흡을 늦추었는가?",
            ACCENT_TEAL,
        ),
        (
            "PASS 03",
            "3차 문장·윤문 퇴고 (Line Edit & Proofreading)",
            "소리 내어 읽으며 군더더기 수식어를 덜어내고 문장의 음악성을 다듬는 단계",
            "✓ 상투적 부사(‘매우, 정말, 문득, 갑자기’)와 번역투(‘~에 다름 아니다, ~로부터’) 제거\n"
            "✓ 한 문장에 중복된 주어(‘그는 ~ 그의 ~’)와 불필요한 피동형(‘~되어졌다’) 교정\n"
            "✓ 육성으로 낭독(Read Aloud)했을 때 숨이 차거나 어색하게 걸리는 대사 수정",
            ACCENT_GOLD,
        ),
    ]

    for i, (pbadge, ptitle, psub, pcheck, pcol) in enumerate(passes):
        py = 1.56 + i * 1.82
        add_rect(slide, 5.42, py, 7.26, 1.68, fill_color=BG_WHITE, line_color=BORDER_WARM, line_pt=1.0, rounded=True, radius=0.04)
        add_rect(slide, 5.42, py, 0.10, 1.68, fill_color=pcol)

        add_rect(slide, 5.68, py + 0.16, 1.05, 0.34, fill_color=pcol, rounded=True, radius=0.18)
        add_textbox(slide, 5.68, py + 0.16, 1.05, 0.34, text=pbadge, size_pt=10.0, bold=True, color=BG_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        add_textbox(slide, 6.86, py + 0.14, 5.6, 0.30, text=ptitle, size_pt=14.0, bold=True, color=INK_NAVY, margin_in=0.01)
        add_textbox(slide, 6.86, py + 0.46, 5.6, 0.24, text=psub, size_pt=10.2, bold=True, color=pcol, margin_in=0.01)

        add_textbox(slide, 5.68, py + 0.72, 6.8, 0.88, text=pcheck, size_pt=10.3, bold=False, color=INK_SLATE, line_spacing=1.26, margin_in=0.01)


def build_slide_10_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_and_footer(
        slide,
        "09  |  THE WRITER'S HABIT & NEXT STEPS",
        "완벽한 첫 문장을 기다리지 말고, 지금 당장 첫 장면을 시작하라",
        "재능보다 강력한 것은 매일 책상 앞에 앉는 루틴과 끝까지 마침표를 찍어본 완주의 근육이다",
        page_num=10,
        dark=True,
    )

    habits = [
        (
            "HABIT 01",
            "영감이 아닌 ‘분량 루틴’으로 써라",
            "영감은 기다리는 자가 아니라 쓰고 있는 자에게 찾아온다.",
            "• 매일 같은 시간, 같은 장소에서 정해진 분량(예: 하루 1,000자~2,000자 또는 45분)을 채우는 시스템 구축\n"
            "• 다음 날 글쓰기가 막막하지 않도록, 오늘 가장 잘 써지는 장면의 문장 중간에서 멈추는 ‘헤밍웨이 기법’ 활용\n"
            "• 하루 1,500자씩 60일이면 단편 3편 또는 경장편 1권 분량(원고지 450매)이 완성된다",
            ACCENT_RUST,
        ),
        (
            "HABIT 02",
            "일상을 채집하는 ‘작가의 수집함’",
            "독창적인 디테일은 책상 위 공상이 아니라 날것의 관찰에서 나온다.",
            "• 지하철에서 들은 기묘한 대화 한 줄, 독특한 직업군의 은어, 계절별 빛과 냄새를 메모장에 즉시 기록\n"
            "• 마음에 드는 소설을 읽을 때는 독자로서 한 번, 작가로서 ‘어떻게 이 감정을 설계했는가’ 분석하며 한 번 더 필사·해부\n"
            "• 막힐 때마다 꺼내 쓸 수 있는 나만의 ‘인물·장면·어휘 창고’를 상시 비축",
            ACCENT_GOLD,
        ),
        (
            "HABIT 03",
            "미완성 대작보다 ‘완성된 단편 한 편’",
            "결말까지 도달해 본 경험만이 작가를 다음 단계로 성장시킨다.",
            "• 처음부터 거대한 10권짜리 세계관에 매몰되기보다, 원고지 70~100매 분량의 단편으로 기승전결 완주 훈련\n"
            "• 첫 초고가 마음에 들지 않는 것은 당신의 재능이 부족해서가 아니라 안목이 높기 때문임을 기억\n"
            "• 오늘 당장 주인공의 이름과 그가 처한 첫 번째 작은 위기 한 줄을 적는 것으로 시작하라",
            ACCENT_TEAL,
        ),
    ]

    for i, (hbadge, htitle, hquote, hbody, hcol) in enumerate(habits):
        hx = 0.65 + i * 4.10
        hy = 1.62
        hw, hh = 3.83, 4.15
        add_rect(slide, hx, hy, hw, hh, fill_color=BG_DARK_CARD, line_color=BORDER_DARK, line_pt=1.0, rounded=True, radius=0.04)
        add_rect(slide, hx, hy, hw, 0.08, fill_color=hcol)

        add_textbox(slide, hx + 0.22, hy + 0.20, 3.4, 0.26, text=hbadge, size_pt=10.0, bold=True, color=hcol, margin_in=0.01)
        add_textbox(slide, hx + 0.22, hy + 0.48, 3.4, 0.36, text=htitle, size_pt=14.5, bold=True, color=INK_LIGHT, margin_in=0.01)

        add_rect(slide, hx + 0.22, hy + 0.94, 3.39, 0.56, fill_color=BG_DARK, line_color=hcol, line_pt=0.8, rounded=True, radius=0.08)
        add_textbox(slide, hx + 0.32, hy + 0.96, 3.19, 0.52, text=f"“{hquote}”",
                    size_pt=10.0, bold=True, color=ACCENT_GOLD if hcol != ACCENT_GOLD else INK_LIGHT, anchor=MSO_ANCHOR.MIDDLE, margin_in=0.01)

        add_textbox(slide, hx + 0.22, hy + 1.66, 3.39, 2.30, text=hbody,
                    size_pt=10.3, bold=False, color=INK_LIGHT_MUTED, line_spacing=1.32, margin_in=0.01)

    add_rect(slide, 0.65, 6.00, 12.03, 0.86, fill_color=ACCENT_RUST, rounded=True, radius=0.06)
    add_textbox(slide, 0.85, 6.04, 11.63, 0.44,
                text="“최악의 초고라도 백지보다는 언제나 위대하다. 백지는 고칠 수 없지만, 쓰인 문장은 반드시 빛나게 다듬을 수 있다.”",
                size_pt=14.0, bold=True, color=INK_LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 0.85, 6.49, 11.63, 0.32,
                text="오늘 당신의 첫 번째 문장을 시작하세요  —  Happy Writing!",
                size_pt=11.0, bold=False, color=SOFT_GOLD_BG, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover_art = generate_cover_art()
    iceberg_art = generate_iceberg_diagram()
    plot_chart = generate_plot_arc_chart()

    build_slide_1_cover(prs, cover_art)
    build_slide_2_pillars(prs)
    build_slide_3_logline(prs)
    build_slide_4_character(prs, iceberg_art)
    build_slide_5_plot_arc(prs, plot_chart)
    build_slide_6_scene_dynamics(prs)
    build_slide_7_pov(prs)
    build_slide_8_show_dont_tell(prs)
    build_slide_9_revision(prs)
    build_slide_10_closing(prs)

    warnings = validate.check_presentation(prs)
    if not warnings:
        print("[ppt_module.validate] 겹침/칸 밖 이탈 없음")

    out_kr = os.path.join(os.path.dirname(os.path.abspath(__file__)), "소설쓰기_가이드.pptx")
    prs.save(out_kr)
    print("Saved:", out_kr)


if __name__ == "__main__":
    main()
