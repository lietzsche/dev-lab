#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Textbox helper that measures text before drawing it.

Two concrete bugs in the legacy build_novel_ppt.py / build_custom_novel_ppt.py
scripts are fixed here:

1. ``run.text = "line one\\nline two"`` does NOT create a line break in
   python-pptx -- it writes a literal ``\\n`` character into a single
   ``<a:t>`` run, which PowerPoint/LibreOffice render as nothing (the two
   lines get mashed together) rather than as two lines. Every multi-line
   string in the legacy scripts (dozens of call sites) is silently affected.
   ``add_textbox`` here splits on "\\n" and emits separate, physically spaced
   text shapes. This also works around browser OOXML viewers that ignore both
   ``<a:br/>`` and DrawingML paragraph spacing.

2. Box heights in the legacy scripts are hand-picked constants with no
   relation to the actual text length / font size / wrap width, so anything
   longer than the author guessed overflows the box and collides with
   whatever is drawn below it. ``add_textbox`` here measures the text via
   ``text_metrics`` and either (a) computes the height itself, or (b), if a
   fixed height/max_height is given and the text doesn't fit, shrinks the
   font size (PowerPoint's own "shrink text on overflow" behavior) down to
   ``min_size_pt`` before giving up and *loudly* warning -- it never just
   lets text silently spill out of its box.
"""

from dataclasses import dataclass

from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

from .theme import FONT_KR, DEFAULT_THEME
from .shapes import set_font
from .text_metrics import _PARAGRAPH_GAP_FACTOR, count_wrapped_lines, measure_block_height_in

_VERT_MARGIN_RATIO = 0.7  # matches the legacy scripts' top/bottom margin convention
# A box that only just fits according to Pillow is fragile: PowerPoint, LibreOffice,
# and browser-based OOXML viewers use different fallback fonts and line metrics.
# Keep deliberate vertical headroom so a two-line heading remains separated even
# when another renderer measures the same glyphs a little taller.
MULTILINE_FIT_HEADROOM = 0.90


def _fit_headroom(text, size_pt, width_in, bold):
    """Reserve extra height only where line-to-line collisions are possible."""
    if count_wrapped_lines(text, size_pt, width_in, bold=bold) > 1:
        return MULTILINE_FIT_HEADROOM
    return 1.0


@dataclass
class TextBoxResult:
    shape: object
    text_frame: object
    left_in: float
    top_in: float
    width_in: float
    height_in: float
    size_pt: float
    overflowed: bool  # True if the text still doesn't fit even at min_size_pt
    shapes: tuple = ()  # one shape per explicit line; ``shape`` is the first


def _shrink_to_fit(text, width_in, avail_height_in, start_size_pt, bold, line_spacing, min_size_pt, step=0.5):
    size = start_size_pt
    while size > min_size_pt:
        target_height_in = avail_height_in * _fit_headroom(text, size, width_in, bold)
        if measure_block_height_in(text, size, width_in, bold=bold, line_spacing=line_spacing) <= target_height_in:
            return size, False
        size -= step
    target_height_in = avail_height_in * _fit_headroom(text, min_size_pt, width_in, bold)
    overflow = measure_block_height_in(
        text, min_size_pt, width_in, bold=bold, line_spacing=line_spacing
    ) > target_height_in
    return min_size_pt, overflow


def _create_text_shape(slide, left, top, width, height, text, size_pt, *,
                       bold, italic, color, align, anchor, line_spacing,
                       margin_left_right, margin_top_bottom, font_name):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margin_left_right)
    tf.margin_right = Inches(margin_left_right)
    tf.margin_top = Inches(margin_top_bottom)
    tf.margin_bottom = Inches(margin_top_bottom)
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    set_font(run, size_pt, bold=bold, italic=italic, color=color, font_name=font_name)
    return shape


def add_textbox(slide, left, top, width, text, size_pt, *,
                 height=None, min_height=0.18, max_height=None,
                 bold=False, italic=False, color=None,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
                 line_spacing=1.25, margin_in=0.06, font_name=FONT_KR,
                 min_size_pt=8.0, autoshrink=True, warn=True):
    """Adds a textbox whose height is derived from the text unless `height`
    is given, and whose font shrinks (rather than overflows) when the text
    doesn't fit `height`/`max_height`. Returns a TextBoxResult so callers
    (e.g. layout.VStack) know the box's *actual* rendered height.
    """
    if color is None:
        color = DEFAULT_THEME.ink_navy

    content_width_in = max(width - 2 * margin_in, 0.1)
    v_margin_in = margin_in * _VERT_MARGIN_RATIO
    overflowed = False

    if height is None:
        needed_text_h = measure_block_height_in(text, size_pt, content_width_in, bold=bold, line_spacing=line_spacing)
        box_h = needed_text_h / _fit_headroom(text, size_pt, content_width_in, bold) + 2 * v_margin_in
        if max_height is not None and box_h > max_height:
            avail = max_height - 2 * v_margin_in
            size_pt, overflowed = _shrink_to_fit(text, content_width_in, avail, size_pt, bold, line_spacing, min_size_pt)
            box_h = max_height
        box_h = max(box_h, min_height)
    else:
        box_h = height
        if autoshrink:
            avail = max(box_h - 2 * v_margin_in, 0.02)
            needed_text_h = measure_block_height_in(text, size_pt, content_width_in, bold=bold, line_spacing=line_spacing)
            target_height = avail * _fit_headroom(text, size_pt, content_width_in, bold)
            if needed_text_h > target_height:
                size_pt, overflowed = _shrink_to_fit(text, content_width_in, avail, size_pt, bold, line_spacing, min_size_pt)

    if overflowed and warn:
        preview = str(text).strip().splitlines()[0][:40] if str(text).strip() else ""
        print(
            f"[ppt_module.textbox] WARNING: text still overflows its box even at "
            f"{min_size_pt}pt -- '{preview}...' in a {width:.2f}x{box_h:.2f}in box at "
            f"({left:.2f}, {top:.2f})"
        )

    segments = str(text).split("\n") if text else [""]
    shapes = []
    if len(segments) == 1:
        shapes.append(_create_text_shape(
            slide, left, top, width, box_h, segments[0], size_pt,
            bold=bold, italic=italic, color=color, align=align, anchor=anchor,
            line_spacing=line_spacing, margin_left_right=margin_in,
            margin_top_bottom=v_margin_in, font_name=font_name,
        ))
    else:
        # Never depend on a viewer's interpretation of line/paragraph spacing.
        # Each explicit line gets its own non-overlapping shape, separated by a
        # physical gap derived from the final font size.
        gap_in = size_pt * _PARAGRAPH_GAP_FACTOR / 72.0
        available = max(box_h - 2 * v_margin_in - gap_in * (len(segments) - 1), 0.02)
        required = [
            measure_block_height_in(seg, size_pt, content_width_in, bold=bold, line_spacing=line_spacing)
            / _fit_headroom(seg, size_pt, content_width_in, bold)
            for seg in segments
        ]
        extra_per_line = max(available - sum(required), 0.0) / len(segments)
        y = top + v_margin_in
        for segment, required_h in zip(segments, required):
            line_h = required_h + extra_per_line
            shapes.append(_create_text_shape(
                slide, left, y, width, line_h, segment, size_pt,
                bold=bold, italic=italic, color=color, align=align,
                anchor=MSO_ANCHOR.TOP, line_spacing=line_spacing,
                margin_left_right=margin_in, margin_top_bottom=0.0,
                font_name=font_name,
            ))
            y += line_h + gap_in

    tb = shapes[0]
    tf = tb.text_frame

    return TextBoxResult(
        shape=tb, text_frame=tf,
        left_in=left, top_in=top, width_in=width, height_in=box_h,
        size_pt=size_pt, overflowed=overflowed, shapes=tuple(shapes),
    )
