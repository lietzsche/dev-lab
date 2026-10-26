import unittest

from pptx import Presentation
from pptx.util import Inches

from ppt_module import add_textbox, validate
from ppt_module.textbox import MULTILINE_FIT_HEADROOM


class TextboxSafetyTests(unittest.TestCase):
    def _slide(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        return prs, prs.slides.add_slide(prs.slide_layouts[6])

    def test_explicit_newline_is_separate_drawingml_paragraph(self):
        _, slide = self._slide()
        result = add_textbox(slide, 1, 1, 4, "첫 줄\n둘째 줄", 18)
        self.assertNotIn("<a:br/>", result.shape.element.xml)
        self.assertEqual([p.text for p in result.text_frame.paragraphs], ["첫 줄", "둘째 줄"])
        self.assertGreater(result.text_frame.paragraphs[0].space_after.pt, 0)

    def test_fixed_multiline_box_keeps_headroom(self):
        _, slide = self._slide()
        result = add_textbox(
            slide, 1, 1, 4, "첫 줄\n둘째 줄", 22,
            height=1.0, margin_in=0.02, min_size_pt=7,
        )
        fit_ratio, line_count = validate._text_fit_ratio(result.shape)
        self.assertEqual(line_count, 2)
        self.assertLessEqual(fit_ratio, MULTILINE_FIT_HEADROOM + 0.035)

    def test_validator_reports_tightly_packed_multiline_text(self):
        _, slide = self._slide()
        result = add_textbox(
            slide, 1, 1, 4, "첫 줄\n둘째 줄", 20,
            height=0.35, margin_in=0.01, min_size_pt=20,
            autoshrink=False, warn=False,
        )
        warnings = validate.check_slide(slide, index=1)
        self.assertTrue(any("usable height" in warning for warning in warnings))


if __name__ == "__main__":
    unittest.main()
